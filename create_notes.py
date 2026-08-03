#!/usr/bin/env python3
"""Generate HTML & CSS Basic Notes as a Word document."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        if level == 1:
            run.font.color.rgb = RGBColor(0x1A, 0x56, 0x8A)
        elif level == 2:
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        else:
            run.font.color.rgb = RGBColor(0x5B, 0x9B, 0xD5)
    return h


def add_para(doc, text, bold=False, italic=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=11, bold=bold)
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.5 * (level + 1))
    for run in p.runs:
        set_run_font(run, size=11)
    return p


def add_code(doc, code_text):
    """Add a monospace code block paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.25)
    # Light gray shading
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    shd.set(qn("w:val"), "clear")
    p.paragraph_format.element.get_or_add_pPr().append(shd)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run("Note: ")
    set_run_font(run, size=11, bold=True, color=(0xC6, 0x59, 0x11))
    run2 = p.add_run(text)
    set_run_font(run2, size=11)
    run2.italic = True
    return p


def add_table_of_contents_like(doc, items):
    for item in items:
        add_bullet(doc, item)


def build_document():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ========== TITLE ==========
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("HTML & CSS Basic Notes")
    set_run_font(run, name="Calibri", size=28, bold=True, color=(0x1A, 0x56, 0x8A))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("A Beginner's Study Guide")
    set_run_font(run, size=14, color=(0x5B, 0x9B, 0xD5))
    run.italic = True

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("For learning the fundamentals of web page structure and styling")
    set_run_font(run, size=10, color=(0x66, 0x66, 0x66))

    doc.add_paragraph()

    # ========== WHAT IS HTML & CSS ==========
    add_heading_styled(doc, "1. Introduction", 1)

    add_heading_styled(doc, "What is HTML?", 2)
    add_para(
        doc,
        "HTML stands for HyperText Markup Language. It is the standard language used to create the structure and content of web pages. HTML uses tags to mark up text, images, links, forms, and other elements so browsers know how to display them.",
    )
    add_bullet(doc, "HyperText — text that can link to other documents")
    add_bullet(doc, "Markup — tags that describe the structure of content")
    add_bullet(doc, "Language — a set of rules for writing web documents")

    add_heading_styled(doc, "What is CSS?", 2)
    add_para(
        doc,
        "CSS stands for Cascading Style Sheets. It controls how HTML elements look on the screen — colors, fonts, spacing, layout, and responsiveness. HTML builds the structure; CSS styles the appearance.",
    )
    add_bullet(doc, "Cascading — styles can override each other based on priority")
    add_bullet(doc, "Style Sheets — separate rules that describe visual design")

    add_heading_styled(doc, "How They Work Together", 2)
    add_para(
        doc,
        "Think of HTML as the skeleton (bones) of a webpage and CSS as the clothes and design (look). A browser reads the HTML file, builds a document tree, then applies CSS rules to decide how each part should appear.",
    )

    # ========== HTML BASICS ==========
    add_heading_styled(doc, "2. HTML Basics", 1)

    add_heading_styled(doc, "HTML Document Structure", 2)
    add_para(doc, "Every HTML page follows a standard skeleton:")
    add_code(
        doc,
        "<!DOCTYPE html>\n"
        "<html lang=\"en\">\n"
        "<head>\n"
        "  <meta charset=\"UTF-8\">\n"
        "  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\">\n"
        "  <title>My First Page</title>\n"
        "</head>\n"
        "<body>\n"
        "  <h1>Hello, World!</h1>\n"
        "  <p>This is my first webpage.</p>\n"
        "</body>\n"
        "</html>",
    )
    add_bullet(doc, "<!DOCTYPE html> — tells the browser this is an HTML5 document")
    add_bullet(doc, "<html> — root element that wraps the entire page")
    add_bullet(doc, "<head> — metadata (title, charset, CSS links) — not visible on the page")
    add_bullet(doc, "<title> — text shown in the browser tab")
    add_bullet(doc, "<body> — all visible content goes here")

    add_heading_styled(doc, "HTML Tags and Elements", 2)
    add_para(
        doc,
        "An HTML element usually has an opening tag, content, and a closing tag:",
    )
    add_code(doc, "<tagname>content</tagname>\n\nExample: <p>This is a paragraph.</p>")
    add_para(doc, "Some elements are self-closing (void elements) and do not need a closing tag:")
    add_code(doc, '<img src="photo.jpg" alt="A photo">\n<br>\n<hr>')

    add_heading_styled(doc, "Attributes", 2)
    add_para(
        doc,
        "Attributes give extra information about an element. They go inside the opening tag:",
    )
    add_code(doc, '<a href="https://example.com" target="_blank">Visit Example</a>')
    add_bullet(doc, "href — the URL the link points to")
    add_bullet(doc, "target=\"_blank\" — opens the link in a new tab")
    add_bullet(doc, "src — source path for images or scripts")
    add_bullet(doc, "alt — alternative text for images (accessibility)")
    add_bullet(doc, "id — unique identifier for one element")
    add_bullet(doc, "class — reusable label for styling multiple elements")

    add_heading_styled(doc, "Common HTML Tags", 2)

    add_heading_styled(doc, "Headings", 3)
    add_para(doc, "There are six heading levels. <h1> is the most important, <h6> is the least:")
    add_code(doc, "<h1>Main Title</h1>\n<h2>Section Title</h2>\n<h3>Subsection</h3>")
    add_note(doc, "Use only one <h1> per page for good SEO and accessibility.")

    add_heading_styled(doc, "Text Content", 3)
    add_code(
        doc,
        "<p>A paragraph of text.</p>\n"
        "<strong>Bold / important text</strong>\n"
        "<em>Italic / emphasized text</em>\n"
        "<br>  <!-- line break -->\n"
        "<hr>  <!-- horizontal rule / divider -->\n"
        "<blockquote>A quoted passage</blockquote>\n"
        "<code>inline code</code>\n"
        "<pre>preformatted text block</pre>",
    )

    add_heading_styled(doc, "Links", 3)
    add_code(
        doc,
        '<a href="https://google.com">Go to Google</a>\n'
        '<a href="about.html">About page (relative link)</a>\n'
        '<a href="#section1">Jump to section on same page</a>\n'
        '<a href="mailto:hello@example.com">Send email</a>',
    )

    add_heading_styled(doc, "Images", 3)
    add_code(doc, '<img src="images/cat.jpg" alt="A cute cat" width="300">')
    add_note(doc, "Always include alt text so screen readers and slow connections can describe the image.")

    add_heading_styled(doc, "Lists", 3)
    add_para(doc, "Unordered list (bullets):")
    add_code(doc, "<ul>\n  <li>Apple</li>\n  <li>Banana</li>\n  <li>Cherry</li>\n</ul>")
    add_para(doc, "Ordered list (numbers):")
    add_code(doc, "<ol>\n  <li>Wake up</li>\n  <li>Brush teeth</li>\n  <li>Study HTML</li>\n</ol>")
    add_para(doc, "Description list:")
    add_code(
        doc,
        "<dl>\n"
        "  <dt>HTML</dt>\n"
        "  <dd>HyperText Markup Language</dd>\n"
        "  <dt>CSS</dt>\n"
        "  <dd>Cascading Style Sheets</dd>\n"
        "</dl>",
    )

    add_heading_styled(doc, "Tables", 3)
    add_code(
        doc,
        "<table>\n"
        "  <thead>\n"
        "    <tr>\n"
        "      <th>Name</th>\n"
        "      <th>Age</th>\n"
        "    </tr>\n"
        "  </thead>\n"
        "  <tbody>\n"
        "    <tr>\n"
        "      <td>Ali</td>\n"
        "      <td>22</td>\n"
        "    </tr>\n"
        "    <tr>\n"
        "      <td>Sara</td>\n"
        "      <td>25</td>\n"
        "    </tr>\n"
        "  </tbody>\n"
        "</table>",
    )
    add_bullet(doc, "<table> — the whole table")
    add_bullet(doc, "<tr> — table row")
    add_bullet(doc, "<th> — header cell")
    add_bullet(doc, "<td> — data cell")

    add_heading_styled(doc, "Forms (Basics)", 3)
    add_para(doc, "Forms collect user input:")
    add_code(
        doc,
        '<form action="/submit" method="post">\n'
        '  <label for="name">Name:</label>\n'
        '  <input type="text" id="name" name="name" placeholder="Your name">\n'
        "\n"
        '  <label for="email">Email:</label>\n'
        '  <input type="email" id="email" name="email">\n'
        "\n"
        '  <label for="pass">Password:</label>\n'
        '  <input type="password" id="pass" name="password">\n'
        "\n"
        '  <label>\n'
        '    <input type="checkbox" name="subscribe"> Subscribe\n'
        "  </label>\n"
        "\n"
        '  <label for="city">City:</label>\n'
        '  <select id="city" name="city">\n'
        '    <option value="delhi">Delhi</option>\n'
        '    <option value="mumbai">Mumbai</option>\n'
        "  </select>\n"
        "\n"
        '  <textarea name="message" rows="4" cols="40"></textarea>\n'
        "\n"
        '  <button type="submit">Submit</button>\n'
        "</form>",
    )
    add_para(doc, "Common input types:")
    add_bullet(doc, "text, email, password, number, date")
    add_bullet(doc, "checkbox, radio, file, submit, reset")
    add_bullet(doc, "color, range, search, tel, url")

    add_heading_styled(doc, "Semantic HTML", 2)
    add_para(
        doc,
        "Semantic tags describe the meaning of content, not just how it looks. They help SEO, accessibility, and code clarity.",
    )
    add_code(
        doc,
        "<header>  <!-- site or page header -->\n"
        "<nav>     <!-- navigation links -->\n"
        "<main>    <!-- main content of the page -->\n"
        "<section> <!-- a thematic group of content -->\n"
        "<article> <!-- independent, self-contained content -->\n"
        "<aside>   <!-- related / sidebar content -->\n"
        "<footer>  <!-- site or page footer -->\n"
        "<figure>  <!-- image with optional caption -->\n"
        "  <img src=\"chart.png\" alt=\"Sales chart\">\n"
        "  <figcaption>Q1 Sales</figcaption>\n"
        "</figure>",
    )
    add_note(
        doc,
        "Prefer semantic tags over many nested <div> elements when the meaning is clear.",
    )

    add_heading_styled(doc, "Block vs Inline Elements", 2)
    add_para(doc, "Block elements start on a new line and take full width:")
    add_bullet(doc, "Examples: <div>, <p>, <h1>–<h6>, <ul>, <section>, <header>")
    add_para(doc, "Inline elements sit within a line of text and only take needed width:")
    add_bullet(doc, "Examples: <span>, <a>, <strong>, <em>, <img>, <code>")

    add_heading_styled(doc, "Comments in HTML", 2)
    add_code(doc, "<!-- This is an HTML comment. Browsers ignore it. -->")

    # ========== CSS BASICS ==========
    add_heading_styled(doc, "3. CSS Basics", 1)

    add_heading_styled(doc, "Ways to Add CSS", 2)
    add_para(doc, "1. Inline CSS — style attribute on an element (avoid for most cases):")
    add_code(doc, '<p style="color: red; font-size: 18px;">Hello</p>')
    add_para(doc, "2. Internal CSS — <style> tag inside <head>:")
    add_code(
        doc,
        "<head>\n"
        "  <style>\n"
        "    p { color: blue; }\n"
        "  </style>\n"
        "</head>",
    )
    add_para(doc, "3. External CSS — separate .css file (best practice):")
    add_code(doc, '<link rel="stylesheet" href="styles.css">')
    add_note(doc, "External CSS is preferred because one file can style many pages and is easier to maintain.")

    add_heading_styled(doc, "CSS Syntax", 2)
    add_code(
        doc,
        "selector {\n"
        "  property: value;\n"
        "  another-property: value;\n"
        "}\n"
        "\n"
        "/* Example */\n"
        "h1 {\n"
        "  color: navy;\n"
        "  font-size: 32px;\n"
        "  text-align: center;\n"
        "}",
    )
    add_bullet(doc, "Selector — which element(s) to style")
    add_bullet(doc, "Property — what to change (color, size, margin, etc.)")
    add_bullet(doc, "Value — how to change it")
    add_bullet(doc, "Each declaration ends with a semicolon (;)")

    add_heading_styled(doc, "CSS Selectors", 2)
    add_para(doc, "Element selector — all matching tags:")
    add_code(doc, "p { color: gray; }")
    add_para(doc, "Class selector — elements with that class (starts with .):")
    add_code(doc, ".highlight { background-color: yellow; }\n\n<p class=\"highlight\">Important text</p>")
    add_para(doc, "ID selector — one unique element (starts with #):")
    add_code(doc, "#hero { font-size: 40px; }\n\n<h1 id=\"hero\">Welcome</h1>")
    add_para(doc, "Universal selector — everything:")
    add_code(doc, "* { box-sizing: border-box; }")
    add_para(doc, "Descendant / child / grouping:")
    add_code(
        doc,
        "/* Any <a> inside <nav> */\n"
        "nav a { text-decoration: none; }\n"
        "\n"
        "/* Direct child only */\n"
        "ul > li { margin: 4px 0; }\n"
        "\n"
        "/* Group multiple selectors */\n"
        "h1, h2, h3 { font-family: Georgia, serif; }",
    )
    add_para(doc, "Attribute and pseudo-class selectors:")
    add_code(
        doc,
        'input[type="email"] { border: 1px solid #ccc; }\n'
        "\n"
        "a:hover { color: orange; }      /* mouse over */\n"
        "a:visited { color: purple; }    /* already clicked */\n"
        "li:first-child { font-weight: bold; }\n"
        "p:nth-child(2) { color: teal; }",
    )

    add_heading_styled(doc, "Colors", 2)
    add_code(
        doc,
        "color: red;                 /* named color */\n"
        "color: #1a568a;             /* hex */\n"
        "color: rgb(26, 86, 138);    /* RGB */\n"
        "color: rgba(26, 86, 138, 0.8); /* RGB + opacity */\n"
        "background-color: #f5f5f5;",
    )

    add_heading_styled(doc, "Typography (Text & Fonts)", 2)
    add_code(
        doc,
        "font-family: Arial, Helvetica, sans-serif;\n"
        "font-size: 16px;            /* also: em, rem, % */\n"
        "font-weight: bold;          /* or 400, 700 */\n"
        "font-style: italic;\n"
        "line-height: 1.5;           /* spacing between lines */\n"
        "letter-spacing: 1px;\n"
        "text-align: center;         /* left | right | center | justify */\n"
        "text-decoration: none;      /* remove underlines on links */\n"
        "text-transform: uppercase;\n"
        "color: #333;",
    )
    add_note(
        doc,
        "rem is relative to the root font size (usually 16px). 1rem = 16px by default. Prefer rem for scalable designs.",
    )

    add_heading_styled(doc, "The Box Model", 2)
    add_para(
        doc,
        "Every HTML element is a box made of four layers (from inside out):",
    )
    add_bullet(doc, "Content — text or images")
    add_bullet(doc, "Padding — space inside the border, around the content")
    add_bullet(doc, "Border — line around the padding")
    add_bullet(doc, "Margin — space outside the border, between elements")
    add_code(
        doc,
        ".card {\n"
        "  width: 300px;\n"
        "  padding: 20px;              /* all sides */\n"
        "  border: 2px solid #333;\n"
        "  margin: 10px 20px;          /* top/bottom | left/right */\n"
        "  box-sizing: border-box;     /* width includes padding + border */\n"
        "}",
    )
    add_para(doc, "Shorthand order for margin/padding (clockwise): top → right → bottom → left")
    add_code(
        doc,
        "margin: 10px;                /* all four sides */\n"
        "margin: 10px 20px;           /* vertical | horizontal */\n"
        "margin: 10px 15px 20px 25px; /* top right bottom left */",
    )
    add_note(
        doc,
        "Always set box-sizing: border-box on * so width calculations include padding and border.",
    )

    add_heading_styled(doc, "Display Property", 2)
    add_code(
        doc,
        "display: block;     /* full width, new line */\n"
        "display: inline;    /* sits in text flow */\n"
        "display: inline-block; /* inline but accepts width/height */\n"
        "display: none;      /* hide completely */\n"
        "display: flex;      /* flexible layout */\n"
        "display: grid;      /* two-dimensional grid layout */",
    )
    add_para(doc, "Visibility vs display:")
    add_code(
        doc,
        "visibility: hidden; /* invisible but still takes space */\n"
        "display: none;      /* removed from layout entirely */",
    )

    add_heading_styled(doc, "Width, Height & Overflow", 2)
    add_code(
        doc,
        "width: 100%;\n"
        "max-width: 800px;\n"
        "min-height: 200px;\n"
        "overflow: auto;     /* scroll if content overflows */\n"
        "overflow: hidden;   /* clip overflowing content */",
    )

    add_heading_styled(doc, "Backgrounds & Borders", 2)
    add_code(
        doc,
        "background-color: #eef;\n"
        "background-image: url(\"bg.jpg\");\n"
        "background-size: cover;\n"
        "background-position: center;\n"
        "background-repeat: no-repeat;\n"
        "\n"
        "border: 1px solid #ccc;\n"
        "border-radius: 8px;   /* rounded corners */\n"
        "border-top: 3px solid navy;\n"
        "\n"
        "box-shadow: 0 2px 8px rgba(0, 0, 0, 0.15);",
    )

    add_heading_styled(doc, "Flexbox (Basic Layout)", 2)
    add_para(
        doc,
        "Flexbox makes it easy to align and distribute items in a row or column:",
    )
    add_code(
        doc,
        ".container {\n"
        "  display: flex;\n"
        "  flex-direction: row;      /* row | column */\n"
        "  justify-content: center;  /* main axis: start | center | space-between | space-around */\n"
        "  align-items: center;      /* cross axis alignment */\n"
        "  gap: 16px;                /* space between items */\n"
        "  flex-wrap: wrap;          /* allow wrapping */\n"
        "}\n"
        "\n"
        ".item {\n"
        "  flex: 1;                  /* grow to fill space */\n"
        "}",
    )

    add_heading_styled(doc, "CSS Grid (Basic Layout)", 2)
    add_code(
        doc,
        ".grid {\n"
        "  display: grid;\n"
        "  grid-template-columns: 1fr 1fr 1fr;  /* 3 equal columns */\n"
        "  gap: 20px;\n"
        "}\n"
        "\n"
        "/* Or: two columns, sidebar + main */\n"
        ".layout {\n"
        "  display: grid;\n"
        "  grid-template-columns: 250px 1fr;\n"
        "  gap: 24px;\n"
        "}",
    )

    add_heading_styled(doc, "Positioning", 2)
    add_code(
        doc,
        "position: static;    /* default flow */\n"
        "position: relative;  /* offset from normal position; creates positioning context */\n"
        "position: absolute;  /* positioned relative to nearest positioned ancestor */\n"
        "position: fixed;     /* fixed to the viewport (stays on scroll) */\n"
        "position: sticky;    /* sticks when scrolling past a point */\n"
        "\n"
        ".badge {\n"
        "  position: absolute;\n"
        "  top: 10px;\n"
        "  right: 10px;\n"
        "}",
    )
    add_para(doc, "z-index controls stacking order (which element appears on top):")
    add_code(doc, "z-index: 10;  /* higher number = on top */")

    add_heading_styled(doc, "Units", 2)
    add_bullet(doc, "px — absolute pixels")
    add_bullet(doc, "% — relative to parent")
    add_bullet(doc, "em — relative to parent font-size")
    add_bullet(doc, "rem — relative to root (html) font-size")
    add_bullet(doc, "vw / vh — 1% of viewport width / height")
    add_bullet(doc, "fr — fraction of free space in CSS Grid")

    add_heading_styled(doc, "Cascading & Specificity", 2)
    add_para(
        doc,
        "When multiple rules target the same element, the browser decides which wins using:",
    )
    add_bullet(doc, "Importance — !important (use rarely)")
    add_bullet(doc, "Specificity — ID > class > element")
    add_bullet(doc, "Source order — later rule wins if specificity is equal")
    add_code(
        doc,
        "/* Specificity examples (rough scores) */\n"
        "p { }           /* element = low */\n"
        ".note { }       /* class = medium */\n"
        "#title { }      /* id = high */\n"
        "\n"
        "/* Later rule wins when equal */\n"
        "p { color: blue; }\n"
        "p { color: green; }  /* green wins */",
    )

    add_heading_styled(doc, "Responsive Design Basics", 2)
    add_para(
        doc,
        "Make pages look good on phones, tablets, and desktops using the viewport meta tag and media queries:",
    )
    add_code(
        doc,
        '<!-- In HTML <head> -->\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1.0">\n'
        "\n"
        "/* In CSS */\n"
        ".container {\n"
        "  width: 90%;\n"
        "  max-width: 1100px;\n"
        "  margin: 0 auto;\n"
        "}\n"
        "\n"
        "/* Styles for screens up to 768px wide */\n"
        "@media (max-width: 768px) {\n"
        "  .container {\n"
        "    width: 95%;\n"
        "  }\n"
        "  nav {\n"
        "    flex-direction: column;\n"
        "  }\n"
        "}",
    )

    add_heading_styled(doc, "Transitions (Simple Motion)", 2)
    add_code(
        doc,
        ".button {\n"
        "  background: navy;\n"
        "  color: white;\n"
        "  transition: background 0.3s ease;\n"
        "}\n"
        "\n"
        ".button:hover {\n"
        "  background: darkorange;\n"
        "}",
    )

    add_heading_styled(doc, "Comments in CSS", 2)
    add_code(doc, "/* This is a CSS comment */")

    # ========== BEST PRACTICES ==========
    add_heading_styled(doc, "4. Best Practices for Beginners", 1)
    add_bullet(doc, "Write clean, indented HTML — nest tags properly and close them.")
    add_bullet(doc, "Use semantic tags (<header>, <main>, <nav>, <footer>) instead of only <div>.")
    add_bullet(doc, "Keep CSS in external files; avoid inline styles for reusable design.")
    add_bullet(doc, "Use class names that describe purpose (e.g. .card-title), not appearance (.big-red).")
    add_bullet(doc, "Always add alt text to images.")
    add_bullet(doc, "Set box-sizing: border-box globally.")
    add_bullet(doc, "Test your page on different screen sizes.")
    add_bullet(doc, "Validate HTML at https://validator.w3.org when unsure.")
    add_bullet(doc, "Practice by building small pages: a profile card, a simple form, a blog layout.")

    # ========== QUICK REFERENCE ==========
    add_heading_styled(doc, "5. Quick Reference Cheatsheet", 1)

    add_heading_styled(doc, "Must-Know HTML Tags", 2)
    tags = [
        ("<!DOCTYPE html>", "Document type declaration"),
        ("<html>, <head>, <body>", "Page structure"),
        ("<title>, <meta>, <link>", "Metadata and CSS link"),
        ("<h1>–<h6>, <p>", "Headings and paragraphs"),
        ("<a>, <img>", "Links and images"),
        ("<ul>, <ol>, <li>", "Lists"),
        ("<table>, <tr>, <th>, <td>", "Tables"),
        ("<form>, <input>, <button>, <label>", "Forms"),
        ("<header>, <nav>, <main>, <section>, <footer>", "Semantic layout"),
        ("<div>, <span>", "Generic containers (block / inline)"),
    ]
    for tag, desc in tags:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(tag)
        set_run_font(r1, name="Consolas", size=10, bold=True)
        r2 = p.add_run(f" — {desc}")
        set_run_font(r2, size=11)

    add_heading_styled(doc, "Must-Know CSS Properties", 2)
    props = [
        ("color, background-color", "Text and background color"),
        ("font-family, font-size, font-weight", "Typography"),
        ("margin, padding, border", "Box model spacing"),
        ("width, height, max-width", "Sizing"),
        ("display, position", "Layout behavior"),
        ("flex / grid properties", "Modern layout systems"),
        ("text-align, text-decoration", "Text alignment and underlines"),
        ("border-radius, box-shadow", "Rounded corners and shadow"),
        ("transition", "Smooth hover effects"),
        ("@media", "Responsive breakpoints"),
    ]
    for prop, desc in props:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(prop)
        set_run_font(r1, name="Consolas", size=10, bold=True)
        r2 = p.add_run(f" — {desc}")
        set_run_font(r2, size=11)

    # ========== MINI PROJECT ==========
    add_heading_styled(doc, "6. Mini Practice Project", 1)
    add_para(
        doc,
        "Build a simple personal profile page with the following requirements:",
    )
    add_bullet(doc, "HTML: header with your name, a short bio paragraph, a photo, a list of skills, and a contact form.")
    add_bullet(doc, "CSS: centered layout (max-width ~700px), nice fonts, spaced sections, styled button with hover effect.")
    add_bullet(doc, "Bonus: make it responsive with a media query for mobile.")
    add_para(
        doc,
        "Suggested file structure:",
        bold=True,
    )
    add_code(
        doc,
        "profile/\n"
        "  index.html\n"
        "  styles.css\n"
        "  images/\n"
        "    me.jpg",
    )

    # ========== GLOSSARY ==========
    add_heading_styled(doc, "7. Glossary", 1)
    glossary = [
        ("Element", "An HTML tag pair and its content, e.g. <p>text</p>."),
        ("Attribute", "Extra info on a tag, e.g. href, src, class, id."),
        ("Selector", "CSS pattern that targets HTML elements to style."),
        ("Property", "A style setting in CSS, e.g. color or margin."),
        ("Box model", "Content + padding + border + margin of an element."),
        ("Cascade", "How CSS rules compete and which one wins."),
        ("Specificity", "A score that decides which CSS rule is stronger."),
        ("Semantic HTML", "Tags that describe meaning, not just layout."),
        ("Responsive", "Design that adapts to different screen sizes."),
        ("Viewport", "The visible area of the page in the browser."),
    ]
    for term, meaning in glossary:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{term}: ")
        set_run_font(r1, size=11, bold=True, color=(0x1A, 0x56, 0x8A))
        r2 = p.add_run(meaning)
        set_run_font(r2, size=11)
        p.paragraph_format.space_after = Pt(4)

    # ========== CLOSING ==========
    doc.add_paragraph()
    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = closing.add_run("— End of Notes —")
    set_run_font(run, size=12, bold=True, color=(0x1A, 0x56, 0x8A))

    tip = doc.add_paragraph()
    tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tip.add_run(
        "Tip: Read one section, type the examples yourself in a code editor, "
        "open the HTML file in a browser, then change values and observe the result."
    )
    set_run_font(run, size=10, color=(0x66, 0x66, 0x66))
    run.italic = True

    out_path = "/workspace/HTML_CSS_Basic_Notes.docx"
    doc.save(out_path)
    print(f"Saved: {out_path}")
    return out_path


if __name__ == "__main__":
    build_document()
