#!/usr/bin/env python3
"""Render the study notes Markdown source into a formatted Word (.docx) document.

Supports the Markdown subset used by the notes in ``notes/``: ATX headings,
paragraphs, bullet/numbered lists, fenced code blocks, pipe tables,
blockquote callouts, horizontal rules, a ``[TOC]`` placeholder, and the inline
markers ``**bold**``, ``*italic*`` and ``` `code` ```.

Usage:
    python3 tools/md_to_docx.py notes/html-css-basics.md "HTML & CSS Basics Notes.docx"
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

# ---------------------------------------------------------------------------
# Theme
# ---------------------------------------------------------------------------

BRAND = RGBColor(0x2B, 0x3F, 0xA8)
BRAND_HEX = "2B3FA8"
HEADING2 = RGBColor(0x1F, 0x29, 0x37)
HEADING3 = RGBColor(0x37, 0x41, 0x51)
MUTED = RGBColor(0x5B, 0x61, 0x6B)
CODE_TEXT = RGBColor(0x14, 0x1B, 0x2D)
INLINE_CODE_TEXT = RGBColor(0xA3, 0x1D, 0x38)

CODE_FILL = "F4F6F8"
CODE_BORDER = "D5DBE1"
INLINE_CODE_FILL = "F0F1F3"
CALLOUT_FILL = "EEF2FF"
TABLE_HEADER_FILL = BRAND_HEX
TABLE_STRIPE_FILL = "F5F7FA"
RULE_COLOR = "C7CDD6"

BODY_FONT = "Calibri"
HEAD_FONT = "Calibri"
CODE_FONT = "Consolas"

BODY_SIZE = Pt(10.5)
CODE_SIZE = Pt(8.5)
INLINE_CODE_SIZE = Pt(9.5)


# ---------------------------------------------------------------------------
# Markdown parsing
# ---------------------------------------------------------------------------

@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    lang: str = ""
    items: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    headers: list = field(default_factory=list)


FENCE_RE = re.compile(r"^```\s*([A-Za-z0-9+#-]*)\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)[-*]\s+(.*)$")
ORDERED_RE = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
QUOTE_RE = re.compile(r"^>\s?(.*)$")
TABLE_SEP_RE = re.compile(r"^\|[\s:|-]+\|$")


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_markdown(text: str) -> list[Block]:
    lines = text.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    i = 0

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()

        if not stripped:
            i += 1
            continue

        fence = FENCE_RE.match(stripped)
        if fence:
            lang = fence.group(1)
            i += 1
            code: list[str] = []
            while i < len(lines) and not FENCE_RE.match(lines[i].strip()):
                code.append(lines[i])
                i += 1
            i += 1  # closing fence
            blocks.append(Block("code", lang=lang, items=code))
            continue

        if stripped == "[TOC]":
            blocks.append(Block("toc"))
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            blocks.append(Block("rule"))
            i += 1
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            blocks.append(Block("heading", text=heading.group(2).strip(),
                                level=len(heading.group(1))))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1].strip()):
            headers = split_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1
            blocks.append(Block("table", headers=headers, rows=rows))
            continue

        if QUOTE_RE.match(stripped):
            parts = []
            while i < len(lines) and QUOTE_RE.match(lines[i].strip()):
                parts.append(QUOTE_RE.match(lines[i].strip()).group(1).strip())
                i += 1
            blocks.append(Block("quote", text=" ".join(p for p in parts if p)))
            continue

        if ORDERED_RE.match(raw):
            items = []
            while i < len(lines) and ORDERED_RE.match(lines[i]):
                match = ORDERED_RE.match(lines[i])
                items.append((len(match.group(1)) // 2, match.group(2).strip()))
                i += 1
            blocks.append(Block("ol", items=items))
            continue

        if BULLET_RE.match(raw):
            items = []
            while i < len(lines) and BULLET_RE.match(lines[i]):
                match = BULLET_RE.match(lines[i])
                items.append((len(match.group(1)) // 2, match.group(2).strip()))
                i += 1
            blocks.append(Block("ul", items=items))
            continue

        para: list[str] = []
        while i < len(lines):
            candidate = lines[i]
            trimmed = candidate.strip()
            if (not trimmed or FENCE_RE.match(trimmed) or HEADING_RE.match(trimmed)
                    or QUOTE_RE.match(trimmed) or BULLET_RE.match(candidate)
                    or ORDERED_RE.match(candidate) or trimmed.startswith("|")
                    or trimmed == "[TOC]" or trimmed in ("---", "***", "___")):
                break
            para.append(trimmed)
            i += 1
        if para:
            blocks.append(Block("para", text=" ".join(para)))

    return blocks


# ---------------------------------------------------------------------------
# Low level docx helpers
# ---------------------------------------------------------------------------

def set_shading(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def shade_paragraph(paragraph, fill: str) -> None:
    set_shading(paragraph._p.get_or_add_pPr(), fill)


def shade_run(run, fill: str) -> None:
    set_shading(run._r.get_or_add_rPr(), fill)


def paragraph_borders(paragraph, **sides) -> None:
    """sides: top/bottom/left/right = (size_eighths, color_hex)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        if side not in sides:
            continue
        size, color = sides[side]
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "4")
        element.set(qn("w:color"), color)
        borders.append(element)


def preserve_space(run) -> None:
    for node in run._r.findall(qn("w:t")):
        node.set(qn("xml:space"), "preserve")


def set_cell_margins(cell, top=60, start=90, bottom=60, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def repeat_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def table_full_width(table) -> None:
    tbl_pr = table._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")
    tbl_pr.append(width)


def add_page_number_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, end):
        run._r.append(node)
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    run.font.name = BODY_FONT


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------

def build_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    h1 = styles["Heading 1"]
    h1.font.name = HEAD_FONT
    h1.font.size = Pt(19)
    h1.font.bold = True
    h1.font.color.rgb = BRAND
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = HEAD_FONT
    h2.font.size = Pt(13.5)
    h2.font.bold = True
    h2.font.color.rgb = HEADING2
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = HEAD_FONT
    h3.font.size = Pt(11.5)
    h3.font.bold = True
    h3.font.color.rgb = HEADING3
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2"):
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = BODY_SIZE
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.12

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = styles["Normal"]
    code.font.name = CODE_FONT
    code.font.size = CODE_SIZE
    code.font.color.rgb = CODE_TEXT
    code.paragraph_format.space_before = Pt(6)
    code.paragraph_format.space_after = Pt(10)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Cm(0.15)
    code.paragraph_format.right_indent = Cm(0.05)
    code.paragraph_format.keep_together = True

    numbered = styles.add_style("Numbered Item", WD_STYLE_TYPE.PARAGRAPH)
    numbered.base_style = styles["Normal"]
    numbered.paragraph_format.left_indent = Cm(0.95)
    numbered.paragraph_format.first_line_indent = Cm(-0.65)
    numbered.paragraph_format.space_after = Pt(3)

    callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    callout.base_style = styles["Normal"]
    callout.paragraph_format.left_indent = Cm(0.3)
    callout.paragraph_format.right_indent = Cm(0.15)
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(10)

    cell = styles.add_style("Table Cell", WD_STYLE_TYPE.PARAGRAPH)
    cell.base_style = styles["Normal"]
    cell.font.size = Pt(9.5)
    cell.paragraph_format.space_after = Pt(0)
    cell.paragraph_format.space_before = Pt(0)
    cell.paragraph_format.line_spacing = 1.05

    toc_part = styles.add_style("TOC Part", WD_STYLE_TYPE.PARAGRAPH)
    toc_part.base_style = styles["Normal"]
    toc_part.font.bold = True
    toc_part.font.color.rgb = BRAND
    toc_part.paragraph_format.space_before = Pt(10)
    toc_part.paragraph_format.space_after = Pt(2)
    toc_part.paragraph_format.keep_with_next = True

    toc_item = styles.add_style("TOC Item", WD_STYLE_TYPE.PARAGRAPH)
    toc_item.base_style = styles["Normal"]
    toc_item.font.size = Pt(10)
    toc_item.paragraph_format.left_indent = Cm(0.7)
    toc_item.paragraph_format.space_after = Pt(1)


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------

INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")

# A code block longer than this is allowed to break across pages, so that a
# single long listing does not leave most of the previous page empty.
CODE_KEEP_TOGETHER_LIMIT = 32


def strip_inline_markers(text: str) -> str:
    return re.sub(r"[`*]", "", text)


class Renderer:
    def __init__(self, document: Document):
        self.doc = document

    def inline(self, paragraph, text: str, bold: bool = False, italic: bool = False,
               size: Pt | None = None, color: RGBColor | None = None):
        for token in INLINE_RE.split(text):
            if not token:
                continue
            if len(token) > 1 and token.startswith("`") and token.endswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = CODE_FONT
                run.font.size = INLINE_CODE_SIZE if size is None else size
                run.font.color.rgb = INLINE_CODE_TEXT
                shade_run(run, INLINE_CODE_FILL)
                preserve_space(run)
            elif token.startswith("**") and token.endswith("**"):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
            elif len(token) > 1 and token.startswith("*") and token.endswith("*"):
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            else:
                run = paragraph.add_run(token)
                preserve_space(run)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if size is not None and run.font.name != CODE_FONT:
                run.font.size = size
            if color is not None and run.font.name != CODE_FONT:
                run.font.color.rgb = color
        return paragraph

    def heading(self, block: Block, is_first_part: bool) -> None:
        if block.level == 2 and not is_first_part:
            self.doc.add_page_break()
        style = {2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}.get(block.level, "Heading 3")
        paragraph = self.doc.add_paragraph(style=style)
        self.inline(paragraph, block.text)
        if block.level == 2:
            paragraph_borders(paragraph, bottom=(12, BRAND_HEX))

    def paragraph(self, block: Block) -> None:
        paragraph = self.doc.add_paragraph()
        self.inline(paragraph, block.text)

    def bullets(self, block: Block) -> None:
        for depth, text in block.items:
            style = "List Bullet" if depth == 0 else "List Bullet 2"
            paragraph = self.doc.add_paragraph(style=style)
            self.inline(paragraph, text)

    def numbers(self, block: Block) -> None:
        for index, (_, text) in enumerate(block.items, start=1):
            paragraph = self.doc.add_paragraph(style="Numbered Item")
            marker = paragraph.add_run(f"{index}.\t")
            marker.bold = True
            self.inline(paragraph, text)

    def code(self, block: Block) -> None:
        lines = block.items
        while lines and not lines[0].strip():
            lines.pop(0)
        while lines and not lines[-1].strip():
            lines.pop()
        paragraph = self.doc.add_paragraph(style="Code Block")
        paragraph.paragraph_format.keep_together = len(lines) <= CODE_KEEP_TOGETHER_LIMIT
        shade_paragraph(paragraph, CODE_FILL)
        paragraph_borders(paragraph,
                          top=(6, CODE_BORDER), bottom=(6, CODE_BORDER),
                          left=(18, BRAND_HEX), right=(6, CODE_BORDER))
        for index, line in enumerate(lines):
            run = paragraph.add_run()
            if index:
                run.add_break()
            run.add_text(line if line.strip() else " ")
            preserve_space(run)

    def quote(self, block: Block) -> None:
        paragraph = self.doc.add_paragraph(style="Callout")
        shade_paragraph(paragraph, CALLOUT_FILL)
        paragraph_borders(paragraph, left=(18, BRAND_HEX))
        self.inline(paragraph, block.text)

    def rule(self) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph_borders(paragraph, bottom=(6, RULE_COLOR))

    def table(self, block: Block) -> None:
        table = self.doc.add_table(rows=1, cols=len(block.headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        table_full_width(table)

        header_cells = table.rows[0].cells
        for cell, text in zip(header_cells, block.headers):
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.style = self.doc.styles["Table Cell"]
            self.inline(paragraph, text, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            for run in paragraph.runs:
                run.bold = True
                if run.font.name == CODE_FONT:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_shading(cell._tc.get_or_add_tcPr(), TABLE_HEADER_FILL)
            set_cell_margins(cell)
        repeat_header_row(table.rows[0])

        for row_index, row_values in enumerate(block.rows):
            cells = table.add_row().cells
            for cell, text in zip(cells, row_values):
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.style = self.doc.styles["Table Cell"]
                self.inline(paragraph, text)
                if row_index % 2 == 1:
                    set_shading(cell._tc.get_or_add_tcPr(), TABLE_STRIPE_FILL)
                set_cell_margins(cell)

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        for run in spacer.runs:
            run.font.size = Pt(2)

    def table_of_contents(self, blocks: list[Block]) -> None:
        heading = self.doc.add_paragraph(style="Heading 1")
        heading.add_run("Contents")
        paragraph_borders(heading, bottom=(12, BRAND_HEX))
        for block in blocks:
            if block.kind != "heading":
                continue
            if block.level == 2:
                paragraph = self.doc.add_paragraph(style="TOC Part")
                paragraph.add_run(strip_inline_markers(block.text))
            elif block.level == 3:
                paragraph = self.doc.add_paragraph(style="TOC Item")
                paragraph.add_run(strip_inline_markers(block.text))


# ---------------------------------------------------------------------------
# Document assembly
# ---------------------------------------------------------------------------

def configure_page(document: Document, title: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.different_first_page_header_footer = True

    text_width = Emu(section.page_width - section.left_margin - section.right_margin)

    # The built-in Footer style carries centre/right tab stops sized for Letter
    # paper, so drop them and define a single right stop at the A4 text width.
    document.styles["Footer"].paragraph_format.tab_stops.clear_all()

    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.clear_all()
    footer.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    label = footer.add_run(f"{title}\t")
    label.font.size = Pt(8.5)
    label.font.color.rgb = MUTED
    label.font.name = BODY_FONT
    add_page_number_field(footer, "PAGE")
    paragraph_borders(footer, top=(4, RULE_COLOR))


def add_cover_page(document: Document, title: str, subtitle: str) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(90)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run(title)
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.name = HEAD_FONT
    run.font.color.rgb = BRAND

    tagline = document.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline.paragraph_format.space_after = Pt(14)
    tag_run = tagline.add_run("Beginner Study Notes & Quick Reference")
    tag_run.font.size = Pt(13)
    tag_run.font.color.rgb = HEADING3
    tag_run.font.name = HEAD_FONT

    divider = document.add_paragraph()
    divider.paragraph_format.left_indent = Cm(3.5)
    divider.paragraph_format.right_indent = Cm(3.5)
    divider.paragraph_format.space_after = Pt(16)
    paragraph_borders(divider, bottom=(8, BRAND_HEX))

    summary = document.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Cm(2.0)
    summary.paragraph_format.right_indent = Cm(2.0)
    summary_run = summary.add_run(subtitle)
    summary_run.font.size = Pt(11)
    summary_run.font.color.rgb = MUTED

    for line in ("Part 1  -  Getting Started",
                 "Part 2  -  HTML Basics",
                 "Part 3  -  CSS Basics",
                 "Part 4  -  Putting It Together"):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(line)
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = HEADING2

    document.add_page_break()


def build(md_path: Path, out_path: Path) -> None:
    blocks = parse_markdown(md_path.read_text(encoding="utf-8"))

    title = "HTML & CSS Basic Notes"
    subtitle = ""
    body: list[Block] = []
    for block in blocks:
        if block.kind == "heading" and block.level == 1 and not body:
            title = block.text
            continue
        if block.kind == "para" and not subtitle and not body:
            subtitle = block.text
            continue
        body.append(block)

    document = Document()
    build_styles(document)
    configure_page(document, title)

    document.core_properties.title = title
    document.core_properties.subject = "HTML and CSS fundamentals"
    document.core_properties.author = "Study Notes"
    document.core_properties.keywords = "HTML, CSS, web development, notes, cheat sheet"
    document.core_properties.comments = subtitle

    add_cover_page(document, title, subtitle)

    renderer = Renderer(document)
    seen_part = False
    for block in body:
        if block.kind == "toc":
            renderer.table_of_contents(body)
        elif block.kind == "heading":
            renderer.heading(block, is_first_part=(block.level == 2 and not seen_part))
            if block.level == 2:
                seen_part = True
        elif block.kind == "para":
            renderer.paragraph(block)
        elif block.kind == "ul":
            renderer.bullets(block)
        elif block.kind == "ol":
            renderer.numbers(block)
        elif block.kind == "code":
            renderer.code(block)
        elif block.kind == "quote":
            renderer.quote(block)
        elif block.kind == "table":
            renderer.table(block)
        elif block.kind == "rule":
            renderer.rule()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))

    counts: dict[str, int] = {}
    for block in body:
        counts[block.kind] = counts.get(block.kind, 0) + 1
    print(f"Wrote {out_path}")
    print("Blocks rendered: " + ", ".join(f"{k}={v}" for k, v in sorted(counts.items())))


def main(argv: list[str]) -> int:
    md_path = Path(argv[1]) if len(argv) > 1 else Path("notes/html-css-basics.md")
    out_path = Path(argv[2]) if len(argv) > 2 else Path("HTML_and_CSS_Basic_Notes.docx")
    if not md_path.exists():
        print(f"Source not found: {md_path}", file=sys.stderr)
        return 1
    build(md_path, out_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
