"""Generate 'HTML & CSS Basic Notes' as a Word document (.docx).

Usage:
    python3 generate_notes.py [output_path]

Requires: python-docx  (pip install python-docx)
"""

import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x4E, 0x79)      # dark blue for headings
CODE_BG = "F2F2F2"                        # light grey shading for code blocks
CODE_COLOR = RGBColor(0x9C, 0x27, 0x54)   # magenta-ish for inline code


def set_shading(paragraph, fill):
    """Apply background shading to a paragraph."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    paragraph.paragraph_format.element.get_or_add_pPr().append(shd)


def add_code_block(doc, code):
    """Add a shaded, monospace code block."""
    for line in code.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.25)
        set_shading(p, CODE_BG)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_rich_paragraph(doc, text, style=None):
    """Add a paragraph where `text` may contain `inline code` in backticks."""
    p = doc.add_paragraph(style=style)
    parts = text.split("`")
    for i, part in enumerate(parts):
        run = p.add_run(part)
        if i % 2 == 1:  # inside backticks -> code formatting
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = CODE_COLOR
    return p


def add_bullets(doc, items):
    for item in items:
        add_rich_paragraph(doc, item, style="List Bullet")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            parts = str(value).split("`")
            for j, part in enumerate(parts):
                run = cells[i].paragraphs[0].add_run(part)
                run.font.size = Pt(9.5)
                if j % 2 == 1:
                    run.font.name = "Consolas"
                    run.font.color.rgb = CODE_COLOR
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def style_headings(doc):
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.color.rgb = ACCENT
        style.font.size = Pt(size)
        style.font.bold = True


def build(path):
    doc = Document()
    style_headings(doc)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # ---------------- Title page ----------------
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("\nHTML & CSS\nBasic Notes")
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = ACCENT

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("A Beginner-Friendly Guide to Building Web Pages")
    run.font.size = Pt(14)
    run.font.italic = True

    doc.add_page_break()

    # ---------------- Table of contents (manual) ----------------
    doc.add_heading("Contents", level=1)
    toc = [
        "Part 1 — HTML",
        "   1. What is HTML?",
        "   2. Basic Structure of an HTML Document",
        "   3. Headings and Paragraphs",
        "   4. Text Formatting Tags",
        "   5. Links and Images",
        "   6. Lists",
        "   7. Tables",
        "   8. Forms and Inputs",
        "   9. Semantic HTML5 Elements",
        "   10. HTML Attributes",
        "Part 2 — CSS",
        "   11. What is CSS?",
        "   12. Three Ways to Add CSS",
        "   13. CSS Selectors",
        "   14. Colors and Units",
        "   15. The Box Model",
        "   16. Typography (Text Styling)",
        "   17. Backgrounds and Borders",
        "   18. Display and Positioning",
        "   19. Flexbox Basics",
        "   20. CSS Grid Basics",
        "   21. Responsive Design and Media Queries",
        "   22. Specificity and the Cascade",
        "Part 3 — Putting It Together",
        "   23. Complete Mini Web Page Example",
        "   24. Best Practices and Tips",
        "   25. Practice Ideas and Resources",
    ]
    for line in toc:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)
        if not line.startswith("   "):
            p.runs[0].bold = True
    doc.add_page_break()

    # ================= PART 1: HTML =================
    doc.add_heading("Part 1 — HTML", level=1)

    doc.add_heading("1. What is HTML?", level=2)
    add_bullets(doc, [
        "HTML stands for HyperText Markup Language.",
        "It is the standard language used to create the structure and content of web pages.",
        "HTML is not a programming language — it is a markup language made of elements (tags).",
        "A browser (Chrome, Firefox, Edge) reads HTML files and displays them as web pages.",
        "HTML files are saved with the `.html` extension (for example `index.html`).",
    ])
    add_rich_paragraph(doc, "An HTML element usually has an opening tag, content, and a closing tag:")
    add_code_block(doc, "<tagname>Content goes here...</tagname>\n\n<p>This is a paragraph.</p>")
    add_rich_paragraph(doc, "Some elements are self-closing (empty elements) such as `<br>`, `<hr>`, and `<img>` — they have no closing tag.")

    doc.add_heading("2. Basic Structure of an HTML Document", level=2)
    add_rich_paragraph(doc, "Every HTML page follows this skeleton:")
    add_code_block(doc, """<!DOCTYPE html>
<html lang="en">
  <head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>My First Page</title>
  </head>
  <body>
    <h1>Hello, World!</h1>
    <p>Welcome to my website.</p>
  </body>
</html>""")
    add_table(doc,
              ["Part", "Purpose"],
              [
                  ("`<!DOCTYPE html>`", "Tells the browser this is an HTML5 document. Always the first line."),
                  ("`<html>`", "The root element that wraps the entire page."),
                  ("`<head>`", "Contains information ABOUT the page (title, character set, styles, links). Not visible on the page."),
                  ("`<title>`", "Text shown in the browser tab and search results."),
                  ("`<meta charset=\"UTF-8\">`", "Sets the character encoding so all symbols display correctly."),
                  ("`<body>`", "Contains everything that is visible on the page."),
              ],
              widths=[2.2, 4.3])

    doc.add_heading("3. Headings and Paragraphs", level=2)
    add_bullets(doc, [
        "There are six heading levels: `<h1>` (most important) to `<h6>` (least important).",
        "Use only ONE `<h1>` per page — it is the main title.",
        "`<p>` defines a paragraph. Browsers automatically add space before and after it.",
        "`<br>` inserts a single line break; `<hr>` draws a horizontal line (thematic break).",
    ])
    add_code_block(doc, """<h1>Main Title</h1>
<h2>Section Title</h2>
<h3>Sub-section Title</h3>
<p>This is a paragraph of text.<br>This is a new line in the same paragraph.</p>
<hr>""")

    doc.add_heading("4. Text Formatting Tags", level=2)
    add_table(doc,
              ["Tag", "Meaning", "Example Output"],
              [
                  ("`<strong>`", "Important text (bold)", "This is important"),
                  ("`<b>`", "Bold (visual only, no meaning)", "Bold text"),
                  ("`<em>`", "Emphasized text (italic)", "Emphasized text"),
                  ("`<i>`", "Italic (visual only)", "Italic text"),
                  ("`<u>`", "Underlined text", "Underlined"),
                  ("`<mark>`", "Highlighted text", "Highlighted"),
                  ("`<small>`", "Smaller text", "Fine print"),
                  ("`<del>`", "Deleted (strikethrough) text", "Old price"),
                  ("`<sub>`", "Subscript", "H2O (2 is lowered)"),
                  ("`<sup>`", "Superscript", "x2 (2 is raised)"),
                  ("`<code>`", "Inline computer code", "console.log()"),
              ],
              widths=[1.3, 2.6, 2.6])
    add_rich_paragraph(doc, "Tip: prefer `<strong>` and `<em>` over `<b>` and `<i>` — they carry meaning for screen readers and search engines.")

    doc.add_heading("5. Links and Images", level=2)
    doc.add_heading("Links — the <a> tag", level=3)
    add_code_block(doc, """<a href="https://www.google.com">Visit Google</a>
<a href="about.html">About page (same website)</a>
<a href="https://example.com" target="_blank">Opens in a new tab</a>
<a href="mailto:someone@example.com">Send an email</a>
<a href="#section2">Jump to a section on this page</a>""")
    add_bullets(doc, [
        "`href` — the destination URL (required).",
        "`target=\"_blank\"` — opens the link in a new tab.",
        "Links to sections use `#id` and require an element with that `id` on the page.",
    ])
    doc.add_heading("Images — the <img> tag", level=3)
    add_code_block(doc, """<img src="photo.jpg" alt="A sunset over the ocean" width="400">""")
    add_bullets(doc, [
        "`src` — path or URL of the image (required).",
        "`alt` — text shown if the image cannot load; also read by screen readers (always include it).",
        "`width` / `height` — size in pixels (better controlled with CSS).",
    ])

    doc.add_heading("6. Lists", level=2)
    add_rich_paragraph(doc, "Unordered list (bullets) — `<ul>`, Ordered list (numbers) — `<ol>`, each item is `<li>`:")
    add_code_block(doc, """<ul>
  <li>Coffee</li>
  <li>Tea</li>
</ul>

<ol>
  <li>Wake up</li>
  <li>Brush teeth</li>
</ol>

<dl>                     <!-- Description list -->
  <dt>HTML</dt>
  <dd>The structure of a web page</dd>
  <dt>CSS</dt>
  <dd>The styling of a web page</dd>
</dl>""")
    add_rich_paragraph(doc, "Lists can be nested by placing a `<ul>` or `<ol>` inside an `<li>`.")

    doc.add_heading("7. Tables", level=2)
    add_code_block(doc, """<table>
  <thead>
    <tr>
      <th>Name</th>
      <th>Age</th>
    </tr>
  </thead>
  <tbody>
    <tr>
      <td>Vishnu</td>
      <td>25</td>
    </tr>
    <tr>
      <td>Anu</td>
      <td>23</td>
    </tr>
  </tbody>
</table>""")
    add_table(doc,
              ["Tag", "Purpose"],
              [
                  ("`<table>`", "Wraps the whole table"),
                  ("`<tr>`", "Table row"),
                  ("`<th>`", "Header cell (bold and centered by default)"),
                  ("`<td>`", "Data cell"),
                  ("`<thead>` / `<tbody>`", "Group header rows and body rows"),
                  ("`colspan` / `rowspan`", "Make a cell span multiple columns / rows"),
              ],
              widths=[2.2, 4.3])

    doc.add_heading("8. Forms and Inputs", level=2)
    add_rich_paragraph(doc, "Forms collect user input and send it to a server.")
    add_code_block(doc, """<form action="/submit" method="post">
  <label for="name">Name:</label>
  <input type="text" id="name" name="name" placeholder="Enter your name" required>

  <label for="email">Email:</label>
  <input type="email" id="email" name="email">

  <label for="pass">Password:</label>
  <input type="password" id="pass" name="password">

  <input type="checkbox" id="agree" name="agree">
  <label for="agree">I agree to the terms</label>

  <input type="radio" id="male" name="gender" value="male">
  <label for="male">Male</label>
  <input type="radio" id="female" name="gender" value="female">
  <label for="female">Female</label>

  <select name="country">
    <option value="in">India</option>
    <option value="us">USA</option>
  </select>

  <textarea name="message" rows="4" placeholder="Your message"></textarea>

  <button type="submit">Submit</button>
</form>""")
    add_table(doc,
              ["Input type", "Used for"],
              [
                  ("`text`", "Single-line text"),
                  ("`email`", "Email address (browser validates format)"),
                  ("`password`", "Hidden characters"),
                  ("`number`", "Numeric input with up/down arrows"),
                  ("`checkbox`", "Select multiple options"),
                  ("`radio`", "Select ONE option from a group (same `name`)"),
                  ("`date`", "Date picker"),
                  ("`file`", "Upload a file"),
                  ("`submit`", "Button that submits the form"),
              ],
              widths=[1.6, 4.9])
    add_rich_paragraph(doc, "Always pair inputs with a `<label>` using `for` + `id` — clicking the label focuses the input and it helps accessibility.")

    doc.add_heading("9. Semantic HTML5 Elements", level=2)
    add_rich_paragraph(doc, "Semantic elements describe their meaning to the browser and developer. Prefer them over generic `<div>`s.")
    add_table(doc,
              ["Element", "Meaning"],
              [
                  ("`<header>`", "Top section of a page or article (logo, title, navigation)"),
                  ("`<nav>`", "Navigation links"),
                  ("`<main>`", "The main unique content of the page (one per page)"),
                  ("`<section>`", "A thematic group of content, usually with a heading"),
                  ("`<article>`", "Self-contained content (blog post, news item, comment)"),
                  ("`<aside>`", "Side content (sidebar, related links)"),
                  ("`<footer>`", "Bottom section (copyright, contact info)"),
                  ("`<figure>` + `<figcaption>`", "Image/diagram with a caption"),
                  ("`<div>`", "Generic container with NO meaning (use when nothing else fits)"),
                  ("`<span>`", "Generic INLINE container for styling a piece of text"),
              ],
              widths=[2.2, 4.3])
    add_code_block(doc, """<body>
  <header>
    <h1>My Blog</h1>
    <nav>
      <a href="/">Home</a>
      <a href="/about">About</a>
    </nav>
  </header>
  <main>
    <article>
      <h2>My First Post</h2>
      <p>Post content...</p>
    </article>
  </main>
  <footer>&copy; 2026 Vishnu</footer>
</body>""")

    doc.add_heading("10. HTML Attributes", level=2)
    add_bullets(doc, [
        "Attributes give extra information about elements and are written inside the opening tag: `name=\"value\"`.",
        "`id` — a unique identifier for ONE element on the page.",
        "`class` — a reusable name shared by MANY elements (main hook for CSS).",
        "`style` — inline CSS (avoid; use stylesheets instead).",
        "`title` — tooltip text shown on hover.",
        "`src`, `href`, `alt` — used by images and links as seen earlier.",
        "`data-*` — custom data attributes, e.g. `data-user-id=\"42\"`.",
    ])
    add_code_block(doc, """<p id="intro" class="highlight large" title="Hover text">
  A paragraph with attributes.
</p>""")

    doc.add_page_break()

    # ================= PART 2: CSS =================
    doc.add_heading("Part 2 — CSS", level=1)

    doc.add_heading("11. What is CSS?", level=2)
    add_bullets(doc, [
        "CSS stands for Cascading Style Sheets.",
        "It controls how HTML elements look: colors, fonts, spacing, layout, animation.",
        "HTML = structure (skeleton), CSS = style (skin/clothes).",
        "A CSS rule = selector + declaration block. Each declaration is `property: value;`.",
    ])
    add_code_block(doc, """selector {
  property: value;
  property: value;
}

/* Example */
h1 {
  color: darkblue;
  font-size: 32px;
}""")
    add_rich_paragraph(doc, "Comments in CSS are written as `/* comment */`.")

    doc.add_heading("12. Three Ways to Add CSS", level=2)
    add_rich_paragraph(doc, "1. External stylesheet (BEST — reusable across pages):")
    add_code_block(doc, """<!-- in the <head> of the HTML file -->
<link rel="stylesheet" href="styles.css">""")
    add_rich_paragraph(doc, "2. Internal stylesheet (inside `<style>` in the `<head>`):")
    add_code_block(doc, """<style>
  p { color: green; }
</style>""")
    add_rich_paragraph(doc, "3. Inline style (avoid — hard to maintain):")
    add_code_block(doc, """<p style="color: red; font-size: 18px;">Red text</p>""")

    doc.add_heading("13. CSS Selectors", level=2)
    add_table(doc,
              ["Selector", "Example", "Selects"],
              [
                  ("Element", "`p { }`", "All `<p>` elements"),
                  ("Class", "`.card { }`", "All elements with `class=\"card\"`"),
                  ("ID", "`#header { }`", "The one element with `id=\"header\"`"),
                  ("Universal", "`* { }`", "Every element"),
                  ("Group", "`h1, h2 { }`", "All `<h1>` and `<h2>` elements"),
                  ("Descendant", "`div p { }`", "`<p>` anywhere inside a `<div>`"),
                  ("Child", "`ul > li { }`", "`<li>` that are direct children of `<ul>`"),
                  ("Adjacent sibling", "`h1 + p { }`", "The first `<p>` right after an `<h1>`"),
                  ("Attribute", "`input[type=\"text\"]`", "Inputs with `type=\"text\"`"),
              ],
              widths=[1.4, 1.8, 3.3])
    doc.add_heading("Pseudo-classes and pseudo-elements", level=3)
    add_code_block(doc, """a:hover      { color: orange; }   /* mouse over a link */
input:focus  { border-color: blue; } /* input being typed in */
li:first-child { font-weight: bold; }
li:nth-child(2n) { background: #eee; } /* every 2nd item */
p::first-line  { font-variant: small-caps; }
p::before      { content: ">> "; }""")

    doc.add_heading("14. Colors and Units", level=2)
    doc.add_heading("Color formats", level=3)
    add_code_block(doc, """color: red;                    /* named color */
color: #ff0000;                /* hex code */
color: #f00;                   /* short hex */
color: rgb(255, 0, 0);         /* red, green, blue (0-255) */
color: rgba(255, 0, 0, 0.5);   /* + alpha (transparency 0-1) */
color: hsl(0, 100%, 50%);      /* hue, saturation, lightness */""")
    doc.add_heading("Common units", level=3)
    add_table(doc,
              ["Unit", "Type", "Meaning"],
              [
                  ("`px`", "Absolute", "Pixels — fixed size"),
                  ("`%`", "Relative", "Percentage of the parent element"),
                  ("`em`", "Relative", "Multiple of the element's own font size"),
                  ("`rem`", "Relative", "Multiple of the ROOT (html) font size — great for consistency"),
                  ("`vw` / `vh`", "Relative", "1% of the viewport (browser window) width / height"),
              ],
              widths=[1.0, 1.2, 4.3])
    add_rich_paragraph(doc, "Rule of thumb: use `rem` for font sizes and spacing, `%`/`vw` for fluid layout widths, `px` for small fixed details like borders.")

    doc.add_heading("15. The Box Model", level=2)
    add_rich_paragraph(doc, "Every HTML element is a rectangular box made of four layers (inside to outside):")
    add_bullets(doc, [
        "Content — the text or image itself (`width` and `height`).",
        "Padding — transparent space between content and border.",
        "Border — the line around padding and content.",
        "Margin — transparent space OUTSIDE the border, separating it from other elements.",
    ])
    add_code_block(doc, """.box {
  width: 300px;
  padding: 20px;            /* all four sides */
  border: 2px solid black;
  margin: 10px auto;        /* 10px top/bottom, auto = center horizontally */
}

/* Shorthand order: top right bottom left (clockwise) */
padding: 10px 20px 10px 20px;
margin: 10px 20px;          /* top/bottom 10px, left/right 20px */""")
    add_rich_paragraph(doc, "Important: by default `width` applies only to the content. Add this rule so padding and border are INCLUDED in the width — nearly every project uses it:")
    add_code_block(doc, """* {
  box-sizing: border-box;
}""")

    doc.add_heading("16. Typography (Text Styling)", level=2)
    add_code_block(doc, """p {
  font-family: Arial, Helvetica, sans-serif; /* fallback list */
  font-size: 1rem;          /* 16px by default */
  font-weight: bold;        /* normal | bold | 100-900 */
  font-style: italic;
  line-height: 1.6;         /* space between lines */
  text-align: center;       /* left | right | center | justify */
  text-decoration: underline; /* none removes link underlines */
  text-transform: uppercase;  /* lowercase | capitalize */
  letter-spacing: 1px;
  color: #333333;
}""")
    add_rich_paragraph(doc, "Google Fonts can be added with a `<link>` tag in the HTML head, then used in `font-family`.")

    doc.add_heading("17. Backgrounds and Borders", level=2)
    add_code_block(doc, """.hero {
  background-color: #f5f5f5;
  background-image: url("bg.jpg");
  background-size: cover;       /* fill the area, cropping if needed */
  background-position: center;
  background-repeat: no-repeat;

  border: 1px solid #ccc;       /* width style color */
  border-radius: 8px;           /* rounded corners; 50% = circle */
  box-shadow: 0 2px 8px rgba(0, 0, 0, 0.2); /* x y blur color */
}""")

    doc.add_heading("18. Display and Positioning", level=2)
    doc.add_heading("The display property", level=3)
    add_table(doc,
              ["Value", "Behavior"],
              [
                  ("`block`", "Starts on a new line, takes full width (div, p, h1...)"),
                  ("`inline`", "Flows within text, ignores width/height (span, a, strong...)"),
                  ("`inline-block`", "Flows inline BUT accepts width/height"),
                  ("`none`", "Removes the element completely (takes no space)"),
                  ("`flex`", "Turns the element into a flex container (see Flexbox)"),
                  ("`grid`", "Turns the element into a grid container (see Grid)"),
              ],
              widths=[1.5, 5.0])
    doc.add_heading("The position property", level=3)
    add_table(doc,
              ["Value", "Behavior"],
              [
                  ("`static`", "Default — normal document flow"),
                  ("`relative`", "Offset from its normal spot using `top/right/bottom/left`; original space is kept"),
                  ("`absolute`", "Removed from flow; positioned relative to the nearest positioned ancestor"),
                  ("`fixed`", "Stays in place relative to the browser window (e.g. sticky navbars)"),
                  ("`sticky`", "Scrolls normally, then sticks when it reaches a given offset"),
              ],
              widths=[1.2, 5.3])
    add_code_block(doc, """.badge {
  position: absolute;
  top: 10px;
  right: 10px;
  z-index: 10;   /* higher number = drawn on top */
}""")

    doc.add_heading("19. Flexbox Basics", level=2)
    add_rich_paragraph(doc, "Flexbox arranges items in ONE dimension (a row OR a column). Apply `display: flex` to the PARENT container.")
    add_code_block(doc, """.container {
  display: flex;
  flex-direction: row;        /* row | column */
  justify-content: center;    /* main axis: flex-start | center |
                                 space-between | space-around | space-evenly */
  align-items: center;        /* cross axis: stretch | flex-start | center */
  gap: 16px;                  /* space between items */
  flex-wrap: wrap;            /* allow items to wrap to next line */
}

.item {
  flex: 1;                    /* grow to share space equally */
}""")
    add_rich_paragraph(doc, "The classic \"center anything\" trick:")
    add_code_block(doc, """.parent {
  display: flex;
  justify-content: center;  /* horizontal */
  align-items: center;      /* vertical */
  height: 100vh;
}""")

    doc.add_heading("20. CSS Grid Basics", level=2)
    add_rich_paragraph(doc, "Grid arranges items in TWO dimensions (rows AND columns) at the same time.")
    add_code_block(doc, """.grid {
  display: grid;
  grid-template-columns: 1fr 1fr 1fr;   /* three equal columns */
  grid-template-columns: repeat(3, 1fr); /* same thing, shorter */
  grid-template-rows: 100px auto;
  gap: 20px;
}

.wide-item {
  grid-column: 1 / 3;   /* span from column line 1 to 3 (2 columns) */
}

/* Responsive cards without media queries */
.cards {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(250px, 1fr));
  gap: 16px;
}""")
    add_rich_paragraph(doc, "`fr` is a fraction of the free space. `1fr 2fr` means the second column is twice as wide as the first.")

    doc.add_heading("21. Responsive Design and Media Queries", level=2)
    add_bullets(doc, [
        "Responsive design = the page adapts to any screen size (phone, tablet, desktop).",
        "Always include the viewport meta tag in the HTML head (shown in section 2).",
        "Use relative units (`%`, `rem`, `vw`) and flexible layouts (Flexbox/Grid).",
        "Media queries apply CSS only when a condition is met (e.g. screen width).",
    ])
    add_code_block(doc, """/* Mobile-first: base styles for small screens, then scale up */
.container { flex-direction: column; }

@media (min-width: 768px) {          /* tablets and larger */
  .container { flex-direction: row; }
}

@media (min-width: 1200px) {         /* desktops */
  .container { max-width: 1140px; margin: 0 auto; }
}

img { max-width: 100%; height: auto; } /* responsive images */""")

    doc.add_heading("22. Specificity and the Cascade", level=2)
    add_rich_paragraph(doc, "When multiple rules target the same element, the browser decides the winner using the cascade:")
    add_bullets(doc, [
        "1. Importance — `!important` beats everything (avoid using it).",
        "2. Specificity — inline style > `#id` > `.class` / attribute / pseudo-class > element.",
        "3. Source order — if specificity is equal, the LAST rule in the file wins.",
    ])
    add_code_block(doc, """p { color: black; }            /* specificity: 1  (element)   */
.note { color: blue; }         /* specificity: 10 (class)     */
#special { color: red; }       /* specificity: 100 (id)       */

/* <p id="special" class="note"> will be RED */""")
    add_rich_paragraph(doc, "Inheritance: text properties like `color`, `font-family`, and `line-height` are inherited by child elements automatically; box properties like `margin`, `padding`, and `border` are not.")

    doc.add_page_break()

    # ================= PART 3 =================
    doc.add_heading("Part 3 — Putting It Together", level=1)

    doc.add_heading("23. Complete Mini Web Page Example", level=2)
    add_rich_paragraph(doc, "Save this as `index.html` and open it in a browser:")
    add_code_block(doc, """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>My Profile Card</title>
  <style>
    * { box-sizing: border-box; margin: 0; padding: 0; }

    body {
      font-family: Arial, sans-serif;
      background: #f0f2f5;
      display: flex;
      justify-content: center;
      align-items: center;
      min-height: 100vh;
    }

    .card {
      background: white;
      width: 320px;
      padding: 24px;
      border-radius: 12px;
      box-shadow: 0 4px 12px rgba(0, 0, 0, 0.15);
      text-align: center;
    }

    .card img {
      width: 96px;
      height: 96px;
      border-radius: 50%;
      object-fit: cover;
    }

    .card h1 { font-size: 1.4rem; margin: 12px 0 4px; }
    .card p  { color: #666; font-size: 0.95rem; }

    .card a {
      display: inline-block;
      margin-top: 16px;
      padding: 10px 24px;
      background: #1f4e79;
      color: white;
      text-decoration: none;
      border-radius: 6px;
    }

    .card a:hover { background: #163a5c; }
  </style>
</head>
<body>
  <div class="card">
    <img src="https://via.placeholder.com/96" alt="Profile photo">
    <h1>Vishnu</h1>
    <p>Software Engineer • Web Developer</p>
    <a href="mailto:vishnu@example.com">Contact Me</a>
  </div>
</body>
</html>""")

    doc.add_heading("24. Best Practices and Tips", level=2)
    add_bullets(doc, [
        "Keep HTML for structure and CSS for style — avoid inline styles.",
        "Use semantic tags (`<header>`, `<main>`, `<footer>`) instead of `<div>` everywhere.",
        "Always add `alt` text to images and `<label>` to form inputs (accessibility).",
        "Add `* { box-sizing: border-box; }` at the top of every stylesheet.",
        "Use classes for styling; reserve `id`s for unique anchors and JavaScript hooks.",
        "Design mobile-first, then enhance for larger screens with `min-width` media queries.",
        "Use the browser DevTools (press F12) to inspect and experiment with any page's HTML and CSS live.",
        "Validate your HTML at validator.w3.org to catch mistakes.",
        "Name classes meaningfully: `.profile-card` is better than `.box1`.",
    ])

    doc.add_heading("25. Practice Ideas and Resources", level=2)
    doc.add_heading("Practice projects (in order of difficulty)", level=3)
    add_bullets(doc, [
        "A personal profile card (like the example in section 23).",
        "A recipe page using headings, lists, and images.",
        "A simple survey form using every input type.",
        "A photo gallery using CSS Grid.",
        "A responsive landing page with a navbar, hero section, and footer.",
        "Clone the layout of a simple site you like (e.g. a blog homepage).",
    ])
    doc.add_heading("Free learning resources", level=3)
    add_bullets(doc, [
        "MDN Web Docs (developer.mozilla.org) — the official reference for HTML and CSS.",
        "web.dev/learn — Google's structured HTML and CSS courses.",
        "freeCodeCamp.org — free interactive Responsive Web Design certification.",
        "CSS-Tricks (css-tricks.com) — famous guides for Flexbox and Grid.",
        "Flexbox Froggy and Grid Garden — games to practice Flexbox and Grid.",
    ])

    doc.save(path)
    print(f"Saved: {path}")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "HTML_CSS_Basic_Notes.docx"
    build(out)
