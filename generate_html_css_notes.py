from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("HTML_and_CSS_Basic_Notes.docx")

NAVY = "17324D"
BLUE = "1E6091"
SKY = "EAF4FB"
TEAL = "168AAD"
GREEN = "EAF7EF"
GOLD = "F4A261"
PALE_GOLD = "FFF5E6"
RED = "C44536"
PALE_RED = "FDEDEC"
GRAY = "5C6770"
LIGHT_GRAY = "F3F5F7"
DARK = "1F2933"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text, *, bold=False, color=DARK, size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE, size=9.5)
        shade(table.rows[0].cells[index], BLUE)
        if widths:
            table.rows[0].cells[index].width = Inches(widths[index])
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_text(cells[index], str(value))
            if row_index % 2 == 1:
                shade(cells[index], LIGHT_GRAY)
            if widths:
                cells[index].width = Inches(widths[index])
        prevent_row_split(table.rows[-1])
    document.add_paragraph()
    return table


def set_repeat_text_on_header_footer(paragraph, text, color=GRAY):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_text(document, text="", *, bold=False, italic=False, color=DARK, size=10.5):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def add_bullets(document, items, level=0):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_numbered(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_code(document, code, label=None):
    if label:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(BLUE)
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, NAVY)
    set_cell_margins(cell, top=150, start=190, bottom=150, end=190)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(code.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(WHITE)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_callout(document, title, text, kind="tip"):
    colors = {
        "tip": (TEAL, SKY),
        "remember": (BLUE, SKY),
        "warning": (RED, PALE_RED),
        "practice": (GOLD, PALE_GOLD),
        "success": ("27864A", GREEN),
    }
    accent, fill = colors[kind]
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    title_run = paragraph.add_run(title.upper())
    title_run.bold = True
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(9)
    title_run.font.color.rgb = RGBColor.from_string(accent)
    body = cell.add_paragraph()
    body.paragraph_format.space_after = Pt(0)
    body_run = body.add_run(text)
    body_run.font.name = "Aptos"
    body_run.font.size = Pt(9.5)
    body_run.font.color.rgb = RGBColor.from_string(DARK)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_section(document, number, title, subtitle):
    document.add_page_break()
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run(f"MODULE {number}")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    heading = document.add_heading(title, level=1)
    heading.paragraph_format.space_after = Pt(6)
    lead = add_text(document, subtitle, italic=True, color=GRAY, size=11)
    lead.paragraph_format.space_after = Pt(16)


def add_quick_check(document, questions):
    document.add_heading("Quick check", level=2)
    for index, question in enumerate(questions, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(f"{index}. {question}")
        run.bold = True


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 30, NAVY),
        ("Subtitle", 13, GRAY),
        ("Heading 1", 23, NAVY),
        ("Heading 2", 16, BLUE),
        ("Heading 3", 12, TEAL),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    set_repeat_text_on_header_footer(section.header.paragraphs[0], "HTML & CSS BASIC NOTES  •  BEGINNER GUIDE")
    add_page_number(section.footer.paragraphs[0])


def create_document():
    document = Document()
    configure_document(document)
    document.core_properties.title = "HTML and CSS Basic Notes"
    document.core_properties.subject = "Beginner study guide for learning HTML and CSS"
    document.core_properties.author = "Prepared for Vishnu"
    document.core_properties.keywords = "HTML, CSS, web development, beginner notes"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(82)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("HTML & CSS")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(36)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Basic Notes for Beginners")
    run.font.name = "Aptos Display"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    line_table = document.add_table(rows=1, cols=3)
    line_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, color in zip(line_table.rows[0].cells, (BLUE, TEAL, GOLD)):
        shade(cell, color)
        cell.height = Inches(0.08)

    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_before = Pt(25)
    intro.paragraph_format.left_indent = Inches(0.65)
    intro.paragraph_format.right_indent = Inches(0.65)
    run = intro.add_run(
        "A practical study guide that explains how web pages are structured with HTML "
        "and styled with CSS—using simple examples, exercises, and a mini project."
    )
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    prepared = document.add_paragraph()
    prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prepared.paragraph_format.space_before = Pt(42)
    run = prepared.add_run("Prepared for VISHNU")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    add_callout(
        document,
        "How to use these notes",
        "Type every example yourself. Change one value at a time, save the file, and refresh the browser. "
        "Learning comes from observing what each line changes.",
        "tip",
    )

    document.add_page_break()
    document.add_heading("Learning roadmap", level=1)
    add_table(
        document,
        ["Step", "Topic", "What you will learn"],
        [
            ("1", "Web foundations", "How browsers, HTML, CSS, files, and URLs fit together"),
            ("2", "HTML basics", "Elements, attributes, document structure, text, links, and images"),
            ("3", "Semantic HTML", "Meaningful page structure, lists, tables, forms, and accessibility"),
            ("4", "CSS basics", "Rules, selectors, colors, units, typography, and the cascade"),
            ("5", "Layout", "Box model, display, positioning, Flexbox, and Grid"),
            ("6", "Responsive design", "Mobile-first styling and media queries"),
            ("7", "Mini project", "Build and improve a complete profile card page"),
        ],
        [0.55, 1.6, 4.9],
    )
    document.add_heading("Recommended practice setup", level=2)
    add_bullets(
        document,
        [
            "A code editor such as Visual Studio Code.",
            "A modern browser such as Chrome, Edge, or Firefox.",
            "One project folder containing index.html and styles.css.",
            "Browser Developer Tools: right-click a page and select Inspect.",
        ],
    )
    add_code(
        document,
        """my-first-website/
├── index.html
├── styles.css
└── images/
    └── profile.jpg""",
        "Suggested folder structure",
    )
    add_callout(
        document,
        "Important",
        "File names are case-sensitive on many web servers. Use lowercase names, avoid spaces, and prefer "
        "hyphens—for example, about-me.html.",
        "remember",
    )

    add_section(
        document,
        1,
        "How the web page works",
        "HTML supplies the content and meaning; CSS controls presentation and layout.",
    )
    document.add_heading("HTML, CSS, and JavaScript", level=2)
    add_table(
        document,
        ["Technology", "Main job", "Simple analogy"],
        [
            ("HTML", "Defines content and structure", "The skeleton and rooms of a house"),
            ("CSS", "Controls colors, spacing, and layout", "The paint, furniture, and decoration"),
            ("JavaScript", "Adds behavior and interactivity", "The switches, doors, and appliances"),
        ],
        [1.1, 3.1, 2.85],
    )
    document.add_heading("What happens when you open a page?", level=2)
    add_numbered(
        document,
        [
            "The browser reads the HTML file and builds a document structure.",
            "It downloads linked resources such as CSS files and images.",
            "It applies CSS rules to matching HTML elements.",
            "It calculates layout and paints the result on the screen.",
        ],
    )
    add_callout(
        document,
        "Key idea",
        "A browser ignores most extra whitespace in HTML. Indentation is for humans, so format code clearly.",
        "remember",
    )
    add_quick_check(
        document,
        [
            "Which language describes the meaning and structure of a web page?",
            "Which language controls the visual appearance?",
            "What browser feature lets you inspect and temporarily edit a page?",
        ],
    )

    add_section(
        document,
        2,
        "HTML fundamentals",
        "Learn the document skeleton, elements, attributes, text content, links, and images.",
    )
    document.add_heading("The basic HTML document", level=2)
    add_code(
        document,
        """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>My First Page</title>
  <link rel="stylesheet" href="styles.css">
</head>
<body>
  <h1>Hello, web!</h1>
  <p>This is my first page.</p>
</body>
</html>""",
    )
    add_table(
        document,
        ["Code", "Purpose"],
        [
            ("<!DOCTYPE html>", "Tells the browser to use modern HTML."),
            ('<html lang="en">', "Wraps the page and declares its main language."),
            ("<head>", "Stores metadata and links that are not normal page content."),
            ('<meta charset="UTF-8">', "Supports a wide range of characters."),
            ('<meta name="viewport" ...>', "Makes the page scale correctly on mobile devices."),
            ("<title>", "Sets the browser-tab title and helps search engines."),
            ("<body>", "Contains everything visible on the page."),
        ],
        [2.3, 4.75],
    )
    document.add_heading("Elements, tags, and attributes", level=2)
    add_text(
        document,
        "An element usually has an opening tag, content, and a closing tag. Attributes add extra information.",
    )
    add_code(
        document,
        """<a href="https://example.com" target="_blank">Visit Example</a>

<!-- a = element name
     href and target = attributes
     Visit Example = content -->""",
    )
    add_callout(
        document,
        "Remember",
        "Some elements are void elements and do not have closing tags, including <img>, <br>, <hr>, <meta>, "
        "and <input>.",
        "remember",
    )
    document.add_heading("Headings and text", level=2)
    add_code(
        document,
        """<h1>Main page title</h1>
<h2>Major section</h2>
<h3>Subsection</h3>

<p>A paragraph of normal text.</p>
<p><strong>Important text</strong> and <em>emphasized text</em>.</p>
<blockquote>A longer quotation from another source.</blockquote>""",
    )
    add_bullets(
        document,
        [
            "Use one clear <h1> for the main page topic in most pages.",
            "Follow heading levels in order; do not choose a level only because of its default size.",
            "Use <strong> for importance and <em> for stress or emphasis.",
            "CSS—not repeated <br> tags—should control visual spacing.",
        ],
    )
    document.add_heading("Links", level=2)
    add_code(
        document,
        """<!-- External page -->
<a href="https://developer.mozilla.org/">Learn on MDN</a>

<!-- Another page in the same site -->
<a href="about.html">About me</a>

<!-- Section on the current page -->
<a href="#contact">Jump to contact</a>
<section id="contact">...</section>

<!-- Email link -->
<a href="mailto:hello@example.com">Email me</a>""",
    )
    document.add_heading("Images", level=2)
    add_code(
        document,
        """<img
  src="images/profile.jpg"
  alt="Vishnu smiling outdoors"
  width="400"
  height="400"
>""",
    )
    add_callout(
        document,
        "Accessibility",
        "The alt attribute describes an image's purpose for screen-reader users and appears when the image "
        "cannot load. Use alt=\"\" only for a purely decorative image.",
        "tip",
    )
    add_quick_check(
        document,
        [
            "What belongs inside <head>, and what belongs inside <body>?",
            "Why should heading levels follow a logical order?",
            "Which attribute provides an image description?",
        ],
    )

    add_section(
        document,
        3,
        "Semantic HTML and common content",
        "Use elements that describe what content means, not merely how it should look.",
    )
    document.add_heading("Semantic page structure", level=2)
    add_code(
        document,
        """<body>
  <header>
    <nav aria-label="Main navigation">...</nav>
  </header>

  <main>
    <section>
      <h2>Skills</h2>
      ...
    </section>
    <article>
      <h2>Latest post</h2>
      ...
    </article>
    <aside>Related links</aside>
  </main>

  <footer>Copyright information</footer>
</body>""",
    )
    add_table(
        document,
        ["Element", "Use it for"],
        [
            ("<header>", "Introductory content for a page or section"),
            ("<nav>", "A group of important navigation links"),
            ("<main>", "The page's unique primary content; normally one per page"),
            ("<section>", "A themed group of content, usually with a heading"),
            ("<article>", "Independent content that could stand alone"),
            ("<aside>", "Related but secondary content"),
            ("<footer>", "Closing information for a page or section"),
            ("<div>", "A generic container when no semantic element fits"),
            ("<span>", "A generic inline wrapper for a small piece of content"),
        ],
        [1.35, 5.7],
    )
    document.add_heading("Lists", level=2)
    add_code(
        document,
        """<!-- Unordered list: order is not important -->
<ul>
  <li>HTML</li>
  <li>CSS</li>
</ul>

<!-- Ordered list: sequence matters -->
<ol>
  <li>Create index.html</li>
  <li>Write the document structure</li>
  <li>Open it in a browser</li>
</ol>""",
    )
    document.add_heading("Tables", level=2)
    add_code(
        document,
        """<table>
  <caption>Weekly study plan</caption>
  <thead>
    <tr>
      <th scope="col">Day</th>
      <th scope="col">Topic</th>
    </tr>
  </thead>
  <tbody>
    <tr>
      <td>Monday</td>
      <td>HTML basics</td>
    </tr>
  </tbody>
</table>""",
    )
    add_callout(
        document,
        "Use tables correctly",
        "Tables are for tabular data—not page layout. <caption> names the table, and <th> identifies headers.",
        "warning",
    )
    document.add_heading("Forms", level=2)
    add_code(
        document,
        """<form action="/contact" method="post">
  <div>
    <label for="name">Name</label>
    <input id="name" name="name" type="text" required>
  </div>

  <div>
    <label for="email">Email</label>
    <input id="email" name="email" type="email" required>
  </div>

  <div>
    <label for="message">Message</label>
    <textarea id="message" name="message" rows="5"></textarea>
  </div>

  <button type="submit">Send message</button>
</form>""",
    )
    add_bullets(
        document,
        [
            "Connect every label to its control using matching for and id values.",
            "The name attribute identifies submitted data.",
            "Use the correct input type so browsers can provide helpful keyboards and validation.",
            "Use a <button>, not a clickable <div>, for an action.",
            "HTML validation helps users, but a server must still validate submitted data.",
        ],
    )
    add_quick_check(
        document,
        [
            "When should you use <section> instead of <div>?",
            "Why does an input need a connected <label>?",
            "Are tables appropriate for positioning an entire page?",
        ],
    )

    add_section(
        document,
        4,
        "CSS fundamentals",
        "Write rules, connect a stylesheet, select elements, and understand how competing styles are resolved.",
    )
    document.add_heading("Anatomy of a CSS rule", level=2)
    add_code(
        document,
        """p {
  color: #243b53;
  font-size: 1rem;
}""",
    )
    add_bullets(
        document,
        [
            "Selector: p chooses the elements to style.",
            "Property: color describes what aspect will change.",
            "Value: #243b53 specifies the chosen color.",
            "Declaration: color: #243b53; is one property-value pair.",
        ],
    )
    document.add_heading("Three ways to add CSS", level=2)
    add_table(
        document,
        ["Method", "Example", "Recommendation"],
        [
            ("External", '<link rel="stylesheet" href="styles.css">', "Best for normal projects"),
            ("Internal", "<style>p { color: blue; }</style>", "Useful for a single-page demo"),
            ("Inline", '<p style="color: blue;">...</p>', "Avoid for routine styling"),
        ],
        [1.0, 3.65, 2.4],
    )
    document.add_heading("Common selectors", level=2)
    add_code(
        document,
        """/* Element selector */
p { line-height: 1.6; }

/* Class selector: reusable */
.card { padding: 1rem; }

/* ID selector: unique page target */
#contact { background-color: #f0f7ff; }

/* Descendant selector */
nav a { color: white; }

/* Direct-child selector */
.menu > li { display: inline-block; }

/* Attribute selector */
input[type="email"] { border-color: steelblue; }

/* Pseudo-class */
a:hover { text-decoration: none; }
button:focus-visible { outline: 3px solid orange; }

/* Pseudo-element */
.note::before { content: "Note: "; font-weight: bold; }""",
    )
    add_callout(
        document,
        "Class versus ID",
        "Prefer classes for styling because they are reusable and easier to override. Use IDs mainly for unique "
        "identifiers, form labels, or page-fragment links.",
        "tip",
    )
    document.add_heading("Cascade, specificity, and inheritance", level=2)
    add_text(
        document,
        "When multiple declarations target the same property, the browser considers importance, specificity, "
        "and source order. More specific selectors usually win; if specificity ties, the later rule wins.",
    )
    add_code(
        document,
        """p { color: navy; }          /* lower specificity */
.intro { color: teal; }      /* wins for class="intro" */

/* If specificity is equal, the later declaration wins */
.message { color: blue; }
.message { color: green; }   /* applied */""",
    )
    add_bullets(
        document,
        [
            "Inherited properties commonly include color, font-family, and line-height.",
            "Layout properties such as margin, padding, and border usually do not inherit.",
            "Avoid !important while learning; it makes the cascade harder to manage.",
            "Keep selectors simple and use a consistent class-naming approach.",
        ],
    )
    add_quick_check(
        document,
        [
            "Which selector targets every element with class=\"card\"?",
            "If two selectors have equal specificity, which declaration wins?",
            "Why is an external stylesheet usually preferred?",
        ],
    )

    add_section(
        document,
        5,
        "Colors, units, and typography",
        "Choose readable colors and size content with units that adapt well.",
    )
    document.add_heading("Color formats", level=2)
    add_code(
        document,
        """body {
  color: #1f2933;                    /* hexadecimal */
  background-color: rgb(245, 247, 250);
}

.overlay {
  background-color: rgb(23 50 77 / 80%); /* color with transparency */
}

.accent {
  color: hsl(195 75% 38%);
}""",
    )
    document.add_heading("Useful CSS units", level=2)
    add_table(
        document,
        ["Unit", "Meaning", "Good use"],
        [
            ("px", "CSS pixel; fixed relative to the CSS reference pixel", "Borders and small details"),
            ("rem", "Relative to the root element's font size", "Font sizes, padding, spacing"),
            ("em", "Relative to the current element's font size", "Component-relative spacing"),
            ("%", "Relative to a related containing size", "Fluid widths"),
            ("vw / vh", "Percentage of viewport width / height", "Viewport-based sizing"),
            ("ch", "Approximate width of the '0' character", "Readable text-line widths"),
        ],
        [0.85, 3.35, 2.85],
    )
    add_callout(
        document,
        "Practical default",
        "Use rem for most typography and spacing, percentages or max-width for fluid containers, and px for "
        "thin borders. Do not set all text in fixed pixels without considering browser zoom and user preferences.",
        "tip",
    )
    document.add_heading("Readable typography", level=2)
    add_code(
        document,
        """body {
  font-family: system-ui, -apple-system, "Segoe UI", sans-serif;
  font-size: 1rem;
  line-height: 1.6;
  color: #1f2933;
}

h1, h2, h3 {
  line-height: 1.2;
}

p {
  max-width: 65ch;
}""",
    )
    add_bullets(
        document,
        [
            "Use sufficient contrast between text and its background.",
            "Avoid long paragraphs spanning the full width of a large screen.",
            "Do not rely on color alone to communicate status or errors.",
            "Underline links in body text, or provide another equally clear visual cue.",
        ],
    )
    add_quick_check(
        document,
        [
            "What does rem measure relative to?",
            "Which property controls space between lines of text?",
            "Why might max-width: 65ch improve reading comfort?",
        ],
    )

    add_section(
        document,
        6,
        "The box model and display",
        "Every visible element is a rectangular box made of content, padding, border, and margin.",
    )
    document.add_heading("The four layers", level=2)
    add_table(
        document,
        ["Layer", "Meaning"],
        [
            ("Content", "The text, image, or other content inside the element"),
            ("Padding", "Space between content and border; background extends through it"),
            ("Border", "Line surrounding the padding and content"),
            ("Margin", "Transparent space outside the border, separating elements"),
        ],
        [1.3, 5.75],
    )
    add_code(
        document,
        """/* Apply this useful rule to every element */
*, *::before, *::after {
  box-sizing: border-box;
}

.card {
  width: 20rem;
  padding: 1.5rem;
  border: 1px solid #cbd5e1;
  margin: 2rem auto;
}""",
    )
    add_callout(
        document,
        "Why border-box?",
        "With box-sizing: border-box, a declared width includes content, padding, and border. This makes element "
        "sizes easier to predict.",
        "remember",
    )
    document.add_heading("Display values", level=2)
    add_table(
        document,
        ["Value", "Behavior", "Typical examples"],
        [
            ("block", "Starts on a new line and generally fills available width", "<div>, <p>, <section>"),
            ("inline", "Flows inside text; width and height do not apply normally", "<span>, <a>, <strong>"),
            ("inline-block", "Flows inline but accepts width and height", "Badges, compact controls"),
            ("none", "Removes the element from layout and accessibility tree", "Conditional content"),
            ("flex", "Creates a one-dimensional flex layout", "Navigation, rows, columns"),
            ("grid", "Creates a two-dimensional grid layout", "Cards, page regions"),
        ],
        [1.05, 3.85, 2.15],
    )
    document.add_heading("Overflow", level=2)
    add_code(
        document,
        """.code-sample {
  max-width: 100%;
  overflow-x: auto;
}""",
    )
    add_quick_check(
        document,
        [
            "Does margin appear inside or outside the border?",
            "What does border-box change about width calculation?",
            "Which display mode is designed for a one-dimensional layout?",
        ],
    )

    add_section(
        document,
        7,
        "Layout with Flexbox and Grid",
        "Use Flexbox for one-dimensional alignment and Grid for rows and columns together.",
    )
    document.add_heading("Flexbox", level=2)
    add_code(
        document,
        """.nav {
  display: flex;
  align-items: center;         /* cross axis */
  justify-content: space-between; /* main axis */
  gap: 1rem;
}

.nav-links {
  display: flex;
  gap: 1rem;
  list-style: none;
  margin: 0;
  padding: 0;
}""",
    )
    add_table(
        document,
        ["Property", "Purpose"],
        [
            ("flex-direction", "Sets the main axis: row, row-reverse, column, or column-reverse"),
            ("justify-content", "Aligns or distributes items along the main axis"),
            ("align-items", "Aligns items along the cross axis"),
            ("gap", "Adds consistent space between flex or grid items"),
            ("flex-wrap", "Allows items to move onto additional lines"),
            ("flex", "Controls a flex item's growth, shrinkage, and starting size"),
        ],
        [1.55, 5.5],
    )
    document.add_heading("CSS Grid", level=2)
    add_code(
        document,
        """.card-grid {
  display: grid;
  grid-template-columns: repeat(3, minmax(0, 1fr));
  gap: 1.5rem;
}

/* Automatically fit as many cards as possible */
.responsive-grid {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(16rem, 1fr));
  gap: 1.5rem;
}""",
    )
    add_bullets(
        document,
        [
            "fr means a fraction of the available grid space.",
            "minmax(16rem, 1fr) gives each column a minimum and flexible maximum.",
            "auto-fit lets the number of columns change with available width.",
            "Use gap instead of adding individual margins between every item.",
        ],
    )
    document.add_heading("Positioning", level=2)
    add_table(
        document,
        ["Value", "Behavior"],
        [
            ("static", "Normal document flow; the default"),
            ("relative", "Stays in flow and can become the containing block for positioned children"),
            ("absolute", "Removed from normal flow; positioned against a suitable ancestor"),
            ("fixed", "Removed from flow and positioned relative to the viewport"),
            ("sticky", "Behaves normally until it reaches a specified scroll position"),
        ],
        [1.2, 5.85],
    )
    add_callout(
        document,
        "Layout rule",
        "Prefer normal flow, Flexbox, and Grid for page layout. Use absolute positioning for overlays and small "
        "positioned details, not as the primary layout system.",
        "warning",
    )
    add_quick_check(
        document,
        [
            "When is Grid a better choice than Flexbox?",
            "Which Flexbox property aligns items along the main axis?",
            "Why can absolute positioning be fragile for full-page layout?",
        ],
    )

    add_section(
        document,
        8,
        "Responsive design",
        "Create layouts that remain useful and readable across phones, tablets, and desktops.",
    )
    document.add_heading("Mobile-first CSS", level=2)
    add_text(
        document,
        "Start with styles that work on small screens. Add enhancements when the viewport has enough room.",
    )
    add_code(
        document,
        """.page {
  width: min(100% - 2rem, 70rem);
  margin-inline: auto;
}

.cards {
  display: grid;
  gap: 1rem;
}

@media (min-width: 48rem) {
  .cards {
    grid-template-columns: repeat(2, 1fr);
  }
}

@media (min-width: 70rem) {
  .cards {
    grid-template-columns: repeat(3, 1fr);
  }
}""",
    )
    document.add_heading("Responsive images", level=2)
    add_code(
        document,
        """img {
  display: block;
  max-width: 100%;
  height: auto;
}""",
    )
    add_bullets(
        document,
        [
            "Do not design only for a few named devices; let content guide breakpoints.",
            "Avoid fixed widths that force horizontal scrolling.",
            "Keep touch targets large enough and leave room between them.",
            "Test narrow, medium, and wide widths using browser Developer Tools.",
            "Zoom the page to 200% and check that content remains usable.",
        ],
    )
    add_callout(
        document,
        "Media queries",
        "A media query applies rules only when its condition is true. Keep the base experience usable without "
        "depending on a particular screen width.",
        "remember",
    )
    add_quick_check(
        document,
        [
            "What does mobile-first mean?",
            "Why is max-width: 100% useful for images?",
            "Should breakpoints be chosen only from popular phone models?",
        ],
    )

    add_section(
        document,
        9,
        "A complete mini project",
        "Build a responsive profile card using semantic HTML and reusable CSS.",
    )
    document.add_heading("Step 1: index.html", level=2)
    add_code(
        document,
        """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Vishnu | Profile Card</title>
  <link rel="stylesheet" href="styles.css">
</head>
<body>
  <main class="page">
    <article class="profile-card">
      <img
        class="profile-card__photo"
        src="images/profile.jpg"
        alt="Portrait of Vishnu"
        width="160"
        height="160"
      >

      <div class="profile-card__content">
        <p class="eyebrow">Software Engineer</p>
        <h1>Vishnu</h1>
        <p>
          I enjoy building reliable software and learning modern web
          development.
        </p>

        <ul class="skills" aria-label="Skills">
          <li>HTML</li>
          <li>CSS</li>
          <li>Problem solving</li>
        </ul>

        <a class="button" href="mailto:hello@example.com">Contact me</a>
      </div>
    </article>
  </main>
</body>
</html>""",
    )
    document.add_heading("Step 2: styles.css", level=2)
    add_code(
        document,
        """*, *::before, *::after {
  box-sizing: border-box;
}

:root {
  color-scheme: light;
  font-family: system-ui, -apple-system, "Segoe UI", sans-serif;
  line-height: 1.6;
  color: #1f2933;
  background: #eaf4fb;
}

body {
  margin: 0;
}

.page {
  min-height: 100vh;
  display: grid;
  place-items: center;
  padding: 1.5rem;
}

.profile-card {
  width: min(100%, 48rem);
  display: grid;
  gap: 1.5rem;
  padding: clamp(1.25rem, 4vw, 2.5rem);
  border-radius: 1rem;
  background: white;
  box-shadow: 0 1rem 2.5rem rgb(23 50 77 / 15%);
}

.profile-card__photo {
  width: 10rem;
  aspect-ratio: 1;
  object-fit: cover;
  border-radius: 50%;
}

.profile-card h1 {
  margin-block: 0.15em;
  color: #17324d;
  line-height: 1.1;
}

.eyebrow {
  margin: 0;
  font-weight: 700;
  color: #168aad;
  text-transform: uppercase;
  letter-spacing: 0.08em;
}

.skills {
  display: flex;
  flex-wrap: wrap;
  gap: 0.5rem;
  padding: 0;
  list-style: none;
}

.skills li {
  padding: 0.3rem 0.75rem;
  border-radius: 999px;
  background: #eaf4fb;
}

.button {
  display: inline-block;
  padding: 0.7rem 1rem;
  border-radius: 0.5rem;
  color: white;
  background: #1e6091;
  font-weight: 700;
  text-decoration: none;
}

.button:hover {
  background: #17324d;
}

.button:focus-visible {
  outline: 3px solid #f4a261;
  outline-offset: 3px;
}

@media (min-width: 40rem) {
  .profile-card {
    grid-template-columns: auto 1fr;
    align-items: center;
  }
}""",
    )
    document.add_heading("Step 3: experiments", level=2)
    add_numbered(
        document,
        [
            "Replace the text and image with your own profile information.",
            "Change the accent colors and observe all affected components.",
            "Add another skill and confirm that the list wraps on a narrow screen.",
            "Temporarily remove the media query and compare the layout.",
            "Inspect the card in Developer Tools and adjust padding live.",
            "Use the keyboard Tab key and confirm that the contact link receives a visible focus outline.",
        ],
    )
    add_callout(
        document,
        "Completion check",
        "The page should have no horizontal scrollbar, remain readable when zoomed, show meaningful alternative "
        "text, and provide a visible keyboard focus indicator.",
        "success",
    )

    add_section(
        document,
        10,
        "Debugging and good habits",
        "Use a repeatable process to find errors and write maintainable code.",
    )
    document.add_heading("When HTML or CSS does not work", level=2)
    add_numbered(
        document,
        [
            "Save every edited file and refresh the browser.",
            "Confirm that file names and paths match exactly, including capitalization.",
            "Open Developer Tools and check the Console and Network panels.",
            "Inspect the target element and see whether the selector matches.",
            "Look for crossed-out declarations, which indicate an overridden rule.",
            "Check for missing braces, semicolons, quotation marks, or closing tags.",
            "Reduce the problem to the smallest example, then add code back gradually.",
            "Validate HTML and test in another modern browser when needed.",
        ],
    )
    document.add_heading("Common mistakes", level=2)
    add_table(
        document,
        ["Mistake", "Better approach"],
        [
            ("Using <br> repeatedly for spacing", "Use margin, padding, gap, or layout properties"),
            ("Styling everything with IDs", "Use reusable classes"),
            ("Skipping image alt text", "Write useful alt text or alt=\"\" for decoration"),
            ("Using placeholders as labels", "Provide persistent, connected <label> elements"),
            ("Fixed widths everywhere", "Use fluid widths, max-width, Flexbox, and Grid"),
            ("Removing focus outlines", "Create a clear :focus-visible style"),
            ("Using div for every element", "Choose semantic HTML where it communicates meaning"),
            ("Adding !important repeatedly", "Simplify selectors and understand the cascade"),
        ],
        [2.65, 4.4],
    )
    document.add_heading("Code style checklist", level=2)
    add_bullets(
        document,
        [
            "Indent nested HTML consistently, usually with two spaces.",
            "Use lowercase element, attribute, class, and file names.",
            "Choose class names that describe a component or purpose.",
            "Group related CSS rules and keep selectors short.",
            "Use comments to explain intent, not obvious syntax.",
            "Remove unused rules and test after small changes.",
        ],
    )

    add_section(
        document,
        11,
        "Quick reference",
        "Use these compact tables while practicing.",
    )
    document.add_heading("Essential HTML", level=2)
    add_table(
        document,
        ["Element", "Purpose"],
        [
            ("<h1>–<h6>", "Headings"),
            ("<p>", "Paragraph"),
            ("<a href=\"\">", "Link"),
            ("<img src=\"\" alt=\"\">", "Image"),
            ("<ul>, <ol>, <li>", "Lists"),
            ("<header>, <main>, <footer>", "Major page regions"),
            ("<nav>", "Navigation links"),
            ("<section>, <article>, <aside>", "Meaningful content regions"),
            ("<form>, <label>, <input>", "Form and controls"),
            ("<button>", "User-triggered action"),
            ("<table>, <tr>, <th>, <td>", "Tabular data"),
        ],
        [2.65, 4.4],
    )
    document.add_heading("Essential CSS", level=2)
    add_table(
        document,
        ["Goal", "Useful properties"],
        [
            ("Text", "color, font-family, font-size, font-weight, line-height, text-align"),
            ("Spacing", "margin, padding, gap"),
            ("Size", "width, max-width, min-height"),
            ("Decoration", "background, border, border-radius, box-shadow"),
            ("Display", "display, visibility, overflow"),
            ("Flexbox", "flex-direction, justify-content, align-items, flex-wrap"),
            ("Grid", "grid-template-columns, grid-template-rows, place-items"),
            ("Position", "position, inset, z-index"),
            ("Responsive", "@media, clamp(), min(), max()"),
        ],
        [1.45, 5.6],
    )
    document.add_heading("Selector reference", level=2)
    add_table(
        document,
        ["Selector", "Matches"],
        [
            ("p", "Every <p> element"),
            (".card", "Every element with class=\"card\""),
            ("#intro", "The element with id=\"intro\""),
            ("nav a", "Links anywhere inside <nav>"),
            (".menu > li", "Direct <li> children of .menu"),
            ("input[type=\"email\"]", "Email inputs"),
            ("a:hover", "A link while the pointer is over it"),
            ("button:focus-visible", "A button with visible keyboard focus"),
            ("p::first-line", "The first rendered line of a paragraph"),
        ],
        [2.2, 4.85],
    )

    add_section(
        document,
        12,
        "Practice plan and answers",
        "Reinforce each idea with small tasks before building larger pages.",
    )
    document.add_heading("Seven-session practice plan", level=2)
    add_table(
        document,
        ["Session", "Build"],
        [
            ("1", "A page with headings, paragraphs, emphasis, and comments"),
            ("2", "A navigation menu with internal and external links plus two images"),
            ("3", "A semantic article with lists, a table, and a contact form"),
            ("4", "A stylesheet using classes, colors, units, and typography"),
            ("5", "Three cards demonstrating the box model and Flexbox"),
            ("6", "A responsive card grid using Grid and media queries"),
            ("7", "The profile-card mini project, then review accessibility and debugging"),
        ],
        [1.1, 5.95],
    )
    document.add_heading("Challenge ideas", level=2)
    add_bullets(
        document,
        [
            "Create a personal portfolio home page with About, Skills, Projects, and Contact sections.",
            "Build a recipe page using a list of ingredients and an ordered list of instructions.",
            "Build a pricing section containing three responsive cards.",
            "Re-create a simple layout from a screenshot using only HTML and CSS.",
            "Add a print stylesheet that hides navigation and simplifies colors.",
        ],
    )
    document.add_heading("Quick-check answers", level=2)
    answer_rows = [
        ("Module 1", "HTML; CSS; Developer Tools / Inspect."),
        ("Module 2", "Metadata in <head>, visible content in <body>; logical structure and accessibility; alt."),
        ("Module 3", "Use <section> for a themed, usually headed region; labels identify controls; no."),
        ("Module 4", ".card; the later declaration; reuse, caching, and separation of concerns."),
        ("Module 5", "The root font size; line-height; it limits text to a comfortable line length."),
        ("Module 6", "Outside; declared width includes padding and border; flex."),
        ("Module 7", "When controlling rows and columns; justify-content; it removes content from normal flow."),
        ("Module 8", "Small-screen base styles first; prevents image overflow; no—use content-driven breakpoints."),
    ]
    add_table(document, ["Module", "Answers"], answer_rows, [1.2, 5.85])
    add_callout(
        document,
        "Next step",
        "Build without copying. Start from a blank index.html, write the document skeleton from memory, and add "
        "one section at a time. Consult this guide only when you get stuck.",
        "practice",
    )

    document.add_heading("Trusted learning resources", level=2)
    add_bullets(
        document,
        [
            "MDN Web Docs — developer.mozilla.org",
            "web.dev Learn HTML and Learn CSS — web.dev/learn",
            "W3C HTML Validator — validator.w3.org",
            "W3C CSS Validator — jigsaw.w3.org/css-validator",
        ],
    )
    final = document.add_paragraph()
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    final.paragraph_format.space_before = Pt(24)
    run = final.add_run("Learn • Build • Inspect • Improve")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = create_document()
    print(f"Created {output}")
